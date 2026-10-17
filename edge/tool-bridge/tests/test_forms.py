"""一起填表 (docs/FORM-COLLECT-PLAN.md) —— 两个员工各一个 CATFISH_HOME, 共用一个假邮筒走完整来回。

全部虚构: 「南岭建设」安全部的检查表, 邮箱都是 example.com。
"""
from __future__ import annotations

import json

import pytest

from catfish_tool_bridge import forms, forms_excel, forms_net, forms_tool


class World:
    """假中央邮筒 + 当前登录的是谁 (切 CATFISH_HOME 和 me())。"""

    def __init__(self, tmp_path, monkeypatch):
        self.tmp, self.mp = tmp_path, monkeypatch
        self.box: list[dict] = []
        self.posted: list[dict] = []
        self.who = ""
        monkeypatch.setattr(forms_net, "me", lambda: self.who)
        monkeypatch.setattr(forms_net, "post", self._post)
        monkeypatch.setattr(forms_net, "take", self._take)

    def be(self, email: str) -> None:
        self.who = email
        self.mp.setenv("CATFISH_HOME", str(self.tmp / email.split("@")[0]))

    def _post(self, to, kind, env, ttl):
        payload = json.dumps(env, ensure_ascii=False)   # 跟真的一样过一遍序列化
        self.posted.append({"from": self.who, "to": to, "kind": kind, "ttl": ttl})
        self.box.append({"from": self.who, "to": to, "kind": kind, "env": json.loads(payload)})
        return len(self.box)

    def _take(self):
        mine = [m for m in self.box if m["to"] == self.who]
        for m in mine:
            self.box.remove(m)
        return [{"from": m["from"], "kind": m["kind"], "env": m["env"], "created_at": None} for m in mine]


@pytest.fixture()
def w(tmp_path, monkeypatch):
    return World(tmp_path, monkeypatch)


FIELDS = [
    {"label": "项目部", "type": "text", "required": True},
    {"label": "检查次数", "type": "number", "required": True},
    {"label": "最高隐患等级", "type": "choice", "options": ["无", "一般", "较大", "重大"]},
    {"label": "最近检查日期", "type": "date"},
    {"label": "说明", "type": "longtext"},
]


def _new_form(w, recipients=("zhang@example.com", "li@example.com"), deadline="2099-12-31"):
    w.be("zhou@example.com")
    out = forms.create({"title": "10 月安全检查", "fields": FIELDS, "deadline": deadline,
                        "recipients": list(recipients) + ["zhou@example.com"], "note": "每个项目部一行"})
    return out["form"]["form_id"]


def test_完整来回_发_收_预填_提交_汇总(w):
    fid = _new_form(w)
    sent = forms.send(fid)
    assert sent["sent"] == ["zhang@example.com", "li@example.com"] and not sent["failed"]
    assert all(p["kind"] == "form_invite" for p in w.posted)
    invite = w.box[0]["env"]
    assert "recipients" not in invite["form"] and invite["form"]["title"] == "10 月安全检查"

    w.be("zhang@example.com")
    s = forms.sync()
    assert s["invite"] == 1 and s["new_invites"][0]["from"] == "zhou@example.com"
    d = forms.save_draft(fid, {"project": "一号项目部", "f2": "3"}, {"project": "邮件《10 月检查安排》"})
    assert d["filled"] >= 1
    bad = forms.submit(fid, {"项目部": "一号项目部", "检查次数": "三次"})
    assert not bad["ok"] and any("检查次数" in p for p in bad["problems"])
    keys = {f["label"]: f["key"] for f in forms.load_incoming(fid)["form"]["fields"]}
    ok = forms.submit(fid, {keys["项目部"]: "一号项目部", keys["检查次数"]: "3", keys["最高隐患等级"]: "一般",
                            keys["最近检查日期"]: "2026-10-15"})
    assert ok["ok"] and not ok["late"]

    w.be("li@example.com")
    forms.sync()
    forms.decline(fid, "这个月没有在建项目")

    w.be("zhou@example.com")
    s = forms.sync()
    assert s["response"] == 2
    t = forms.responses_table(fid)
    assert t["rows"] == [{"填写人": "zhang@example.com", "提交时间": t["rows"][0]["提交时间"], "逾期": False,
                          "项目部": "一号项目部", "检查次数": 3, "最高隐患等级": "一般",
                          "最近检查日期": "2026-10-15", "说明": None}]
    assert t["progress"]["counts"] == {"已提交": 1, "不填": 1}
    reasons = {r["email"]: r["reason"] for r in t["progress"]["rows"]}
    assert reasons["li@example.com"] == "这个月没有在建项目"


def test_重复提交覆盖_截止后标逾期(w):
    fid = _new_form(w, recipients=("zhang@example.com",), deadline="2020-01-01")
    forms.send(fid)
    w.be("zhang@example.com")
    forms.sync()
    keys = {f["label"]: f["key"] for f in forms.load_incoming(fid)["form"]["fields"]}
    forms.submit(fid, {keys["项目部"]: "旧值", keys["检查次数"]: 1})
    r = forms.submit(fid, {keys["项目部"]: "新值", keys["检查次数"]: 2})
    assert r["late"]
    w.be("zhou@example.com")
    forms.sync()
    t = forms.responses_table(fid)
    assert [row["项目部"] for row in t["rows"]] == ["新值"] and t["rows"][0]["逾期"]
    assert t["progress"]["counts"] == {"逾期": 1}


def test_催填_结束_收件人那边看得到(w):
    fid = _new_form(w)
    forms.send(fid)
    w.be("zhang@example.com")
    forms.sync()
    w.be("zhou@example.com")
    r = forms.remind(fid, "周五前交哈")
    assert sorted(r["reminded"]) == ["li@example.com", "zhang@example.com"]
    forms.close(fid)
    w.be("zhang@example.com")
    s = forms.sync()
    assert s["remind"] == 1 and s["close"] == 1
    rec = forms.load_incoming(fid)
    assert rec["remind_text"] == "周五前交哈" and rec["status"] == "closed"
    with pytest.raises(ValueError):
        forms.submit(fid, {})


def test_不是收件人的回答不收_冒名的邀请不覆盖(w):
    fid = _new_form(w, recipients=("zhang@example.com",))
    forms.send(fid)
    w.be("mallory@example.com")
    w.box.append({"from": "mallory@example.com", "to": "zhou@example.com", "kind": "form_response",
                  "env": {"v": 1, "form_id": fid, "status": "submitted", "values": {"project": "伪造"}}})
    w.be("zhang@example.com")
    forms.sync()
    w.box.append({"from": "mallory@example.com", "to": "zhang@example.com", "kind": "form_invite",
                  "env": {"v": 1, "form_id": fid, "form": {"title": "钓鱼表", "fields": ["密码"]}}})
    forms.sync()
    assert forms.load_incoming(fid)["form"]["title"] == "10 月安全检查"
    w.be("zhou@example.com")
    assert forms.sync()["response"] == 0
    assert forms.responses_table(fid)["rows"] == []


def test_期限_截止后7天_最长30天_回填至少3天():
    from datetime import datetime, timezone
    now = datetime(2026, 10, 17, tzinfo=timezone.utc)
    day = 24 * 3600
    assert forms_net.ttl_seconds("2026-10-20", now=now) == 11 * day
    assert forms_net.ttl_seconds("2027-06-01", now=now) == 30 * day
    assert forms_net.ttl_seconds(None, now=now) == 30 * day
    assert forms_net.ttl_seconds("2026-01-01", now=now, minimum=3 * day) == 3 * day


def test_字段校验():
    with pytest.raises(ValueError):
        forms.normalize_fields([])
    with pytest.raises(ValueError):
        forms.normalize_fields([{"label": "等级", "type": "choice", "options": ["只有一个"]}])
    with pytest.raises(ValueError):
        forms.normalize_fields([{"label": "x", "type": "formula"}])
    f = forms.normalize_fields(["项目部", "项目部", {"label": "多选", "type": "multichoice", "options": ["a", "b", "c"]}])
    assert len({x["key"] for x in f}) == 3
    clean, problems = forms.check_values(f, {f[2]["key"]: ["c", "a"]}, strict=False)
    assert clean[f[2]["key"]] == ["a", "c"] and not problems


def test_工具_发催结束要员工确认_没有提交动作(w):
    fid = _new_form(w)
    p = forms_tool.dispatch("catfish_form", {"action": "send", "form_id": fid})
    assert p["needs_confirmation"] and "zhang@example.com" in p["proposal"] and w.posted == []
    r = forms_tool.dispatch("catfish_form", {"action": "send", "form_id": fid, "confirmed": True})
    assert r["ok"] and len(w.posted) == 2
    assert forms_tool.dispatch("catfish_form", {"action": "submit", "form_id": fid})["ok"] is False
    from catfish_tool_bridge.catfish_tool_schemas import CATFISH_NATIVE_TOOLS
    schema = next(t for t in CATFISH_NATIVE_TOOLS if t["name"] == "catfish_form")
    assert not {"submit", "decline"} & set(schema["input_schema"]["properties"]["action"]["enum"])


def test_走catfish_tools分派(w):
    from catfish_tool_bridge.catfish_tools import dispatch_native, is_native
    w.be("zhou@example.com")
    assert is_native("catfish_form")
    assert dispatch_native("catfish_form", {"action": "list"}) == {"ok": True, "incoming": [], "outgoing": []}


def _template(path):
    from datetime import date
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(["项目部", "检查次数", "隐患等级", "检查日期", "整改情况"], 1):
        ws.cell(row=1, column=c, value=h)
    ws.cell(row=2, column=4, value=date(2026, 1, 1)).number_format = "yyyy-mm-dd"
    ws.cell(row=2, column=4).value = None
    dv = DataValidation(type="list", formula1='"无,一般,较大,重大"')
    ws.add_data_validation(dv)
    dv.add("C2:C200")
    wb.save(path)


def test_从Excel建表_推断类型_按模板导出(w, tmp_path):
    tpl = tmp_path / "安全检查模板.xlsx"
    _template(tpl)
    got = forms_excel.infer_fields(str(tpl))
    types = {f["label"]: f["type"] for f in got["fields"]}
    assert types == {"项目部": "text", "检查次数": "number", "隐患等级": "choice", "检查日期": "date", "整改情况": "longtext"}
    assert next(f for f in got["fields"] if f["label"] == "隐患等级")["options"] == ["无", "一般", "较大", "重大"]

    w.be("zhou@example.com")
    fid = forms.create({"title": got["title"], "fields": got["fields"], "recipients": ["zhang@example.com"],
                        "template_path": str(tpl)})["form"]["form_id"]
    forms.send(fid)
    w.be("zhang@example.com")
    forms.sync()
    keys = {f["label"]: f["key"] for f in forms.load_incoming(fid)["form"]["fields"]}
    forms.submit(fid, {keys["项目部"]: "二号项目部", keys["检查次数"]: 4, keys["隐患等级"]: "较大",
                       keys["检查日期"]: "2026-10-16", keys["整改情况"]: "已整改"})
    w.be("zhou@example.com")
    forms.sync()
    out = forms_excel.export(fid)
    assert out["rows"] == 1
    from openpyxl import load_workbook
    ws = load_workbook(out["path"]).worksheets[0]
    assert [ws.cell(row=1, column=c).value for c in range(1, 9)] == \
        ["项目部", "检查次数", "隐患等级", "检查日期", "整改情况", "填写人", "提交时间", "状态"]
    assert [ws.cell(row=2, column=c).value for c in (1, 2, 3, 5, 6, 8)] == \
        ["二号项目部", 4, "较大", "已整改", "zhang@example.com", "已提交"]
    assert load_workbook(out["path"])["进度"].cell(row=2, column=2).value == "已提交"


def test_没模板也能导出(w):
    fid = _new_form(w, recipients=("zhang@example.com",))
    forms.send(fid)
    out = forms_excel.export(fid)
    from openpyxl import load_workbook
    wb = load_workbook(out["path"])
    assert wb.worksheets[0].cell(row=1, column=1).value == "项目部" and out["rows"] == 0
    assert wb["进度"].cell(row=2, column=2).value == "未提交"


def test_表单编号防穿越():
    with pytest.raises(ValueError):
        forms.load_incoming("../../etc/passwd")
