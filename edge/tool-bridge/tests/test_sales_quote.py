"""销售跟进阶段 2: 价格表 (sales_pricebook) + 报价单 (sales_quote) + 销售周报 (sales_weekly)。

全部虚构客户 / 虚构商品, 临时 CATFISH_HOME, 不碰员工真知识库。
"""
from __future__ import annotations

import asyncio
import json
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from catfish_tool_bridge import customer_radar, customers, sales_pricebook, sales_quote


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("CATFISH_HOME", str(tmp_path))
    return tmp_path


def c(**args):
    return customers.dispatch("catfish_customer", args)


def make(name="南岭建设", stage="需求确认", **extra):
    extra.setdefault("contacts", [{"name": "周工", "email": "zhou@nanling-build.cn"}])
    assert c(action="upsert", name=name, stage=stage, confirmed=True, owner="林销售", **extra)["ok"]


def xlsx(path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


def pricebook(home: Path) -> dict:
    p = xlsx(home / "价格表.xlsx", [
        ["鲶鱼 2026 价格表"],  # 标题行, 表头在第 2 行
        ["产品名称", "规格", "单位", "单价（元）", "最低价", "备注"],
        ["鲶鱼企业版", "50 人", "套/年", "¥120,000", 96000, "含一年维护"],
        ["鲶鱼企业版", "100 人", "套/年", "20万", 160000, ""],
        ["实施服务", "", "人天", 3000, None, ""],
        ["待定商品", "", "", "面议", None, ""],  # 没价: 跳过
    ])
    return sales_pricebook.import_pricebook(str(p))


# ── 人民币大写 ──

@pytest.mark.parametrize("amount,text", [
    (0, "零元整"), (10, "壹拾元整"), (1001, "壹仟零壹元整"), (12345.6, "壹万贰仟叁佰肆拾伍元陆角整"),
    (100000, "壹拾万元整"), (100010, "壹拾万零壹拾元整"), (10000001, "壹仟万零壹元整"),
    (120000000, "壹亿贰仟万元整"), (1.05, "壹元零伍分"), (0.05, "伍分"), (226000, "贰拾贰万陆仟元整"),
])
def test_rmb_upper(amount, text):
    assert sales_quote.rmb_upper(amount) == text


# ── 价格表导入 ──

def test_import_pricebook_finds_header_and_parses_prices(home):
    r = pricebook(home)
    assert r["count"] == 3 and r["skipped"] == 1
    assert r["columns"]["price"] == "单价(元)" and r["columns"]["name"] == "产品名称"
    ent = [x for x in sales_pricebook.load_pricebook()["items"] if x["spec"] == "100 人"][0]
    assert ent["price"] == 200000 and ent["floor"] == 160000
    assert (home / "sales" / "pricebook_source.xlsx").is_file()
    assert [x["spec"] for x in c(action="price", text="企业版")["items"]] == ["50 人", "100 人"]


def test_import_pricebook_rejects_unknown_headers_and_bad_files(home):
    bad = xlsx(home / "乱表.xlsx", [["东西", "钱"], ["A", 1]])
    with pytest.raises(ValueError, match="东西"):
        sales_pricebook.import_pricebook(str(bad))
    with pytest.raises(ValueError, match=".xlsx"):
        sales_pricebook.import_pricebook(str(home / "老表.xls"))
    broken = home / "坏的.xlsx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="打不开"):
        sales_pricebook.import_pricebook(str(broken))


def test_price_action_without_pricebook_tells_where_to_import(home):
    r = c(action="price", text="企业版")
    assert not r["ok"] and "销售跟进" in r["error"]


# ── 报价单 ──

def test_quote_builtin_versions_and_timeline(home):
    pricebook(home)
    make()
    items = [{"name": "鲶鱼企业版", "spec": "50 人", "qty": 1}, {"name": "实施服务", "qty": 10, "discount": 0.9}]
    r = c(action="quote", customer="南岭建设", items=items, note="合同签订后 2 周内部署")
    assert r["ok"] and r["version"] == 1 and r["total"] == 120000 + 27000
    assert r["total_upper"] == "壹拾肆万柒仟元整"
    ws = load_workbook(r["file"]).active
    cells = [v for row in ws.iter_rows(values_only=True) for v in row if v not in (None, "")]
    assert "南岭建设" in cells and "周工" in cells and 147000 in cells and "9折" in cells
    assert Path(r["file"]).name == "报价-南岭建设-v1.xlsx"

    r2 = c(action="quote", customer="南岭建设", items=items[:1])
    assert r2["version"] == 2 and r2["file"].endswith("-v2.xlsx")
    assert [x["version"] for x in c(action="quotes", customer="南岭建设")["quotes"]] == [2, 1]
    top = c(action="get", customer="南岭建设")["customer"]["timeline"][0]
    assert top["channel"] == "系统" and "v2" in top["text"] and "未发出" in top["text"]


def test_quote_below_floor_needs_confirmation(home):
    pricebook(home)
    make()
    items = [{"name": "鲶鱼企业版", "spec": "50 人", "qty": 1, "discount": 0.7}]  # 84000 < 最低价 96000
    r = c(action="quote", customer="南岭建设", items=items)
    assert r["needs_confirmation"] and r["below_floor"][0]["floor"] == 96000
    assert sales_quote.load_ledger() == []
    r = c(action="quote", customer="南岭建设", items=items, confirmed=True)
    assert r["ok"] and r["total"] == 84000 and r["below_floor_confirmed"]


def test_quote_never_invents_prices(home):
    pricebook(home)
    make()
    r = c(action="quote", customer="南岭建设", items=[{"name": "企业版", "qty": 1}])
    assert not r["ok"] and "鲶鱼企业版 / 50 人" in r["error"]  # 模糊名: 给候选, 不猜
    r = c(action="quote", customer="南岭建设", items=[{"name": "定制接口开发", "qty": 1}])
    assert not r["ok"] and "price" in r["error"]
    r = c(action="quote", customer="南岭建设", items=[{"name": "定制接口开发", "qty": 1, "price": 15000}])
    assert r["ok"] and r["lines"][0]["in_pricebook"] is False
    assert not c(action="quote", customer="科瑞认证", items=[{"name": "实施服务", "qty": 1}])["ok"]  # 没建档的客户


def test_quote_with_company_template(home):
    pricebook(home)
    make()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "报价单 {{报价编号}}"
    ws["A2"], ws["B2"] = "致:", "{{客户}}"
    ws.append(["序号", "商品", "数量", "单价", "金额"])
    ws.append(["{{序号}}", "{{商品}} {{规格}}", "{{数量}}", "{{单价}}", "{{金额}}"])  # 第 4 行 = 明细行
    ws["A6"] = "合计: {{合计大写}}"
    ws["E6"] = "{{合计}}"
    ws.merge_cells("A6:D6")
    tpl = home / "公司模板.xlsx"
    wb.save(tpl)
    assert "商品" in sales_pricebook.import_quote_template(str(tpl))["placeholders"]

    items = [{"name": "实施服务", "qty": 2}, {"name": "实施服务", "qty": 3, "note": "二期"},
             {"name": "鲶鱼企业版", "spec": "100 人", "qty": 1}]
    r = c(action="quote", customer="南岭建设", items=items)
    out = load_workbook(r["file"]).active
    assert out["B2"].value == "南岭建设" and out["A1"].value.startswith("报价单 Q")
    assert [out.cell(row, 2).value for row in (4, 5, 6)] == ["实施服务", "实施服务", "鲶鱼企业版 100 人"]
    assert out["E5"].value == 9000 and out["E6"].value == 200000 and out["A8"].value == "合计: 贰拾壹万伍仟元整" and out["E8"].value == 215000
    assert "A8:D8" in [str(m) for m in out.merged_cells.ranges]  # 合并单元格跟着下移

    sales_pricebook.clear_quote_template()
    assert sales_pricebook.load_quote_settings()["custom_template"] is False


def test_template_without_item_row_is_rejected(home):
    wb = Workbook()
    wb.active["A1"] = "{{客户}}"
    wb.save(home / "t.xlsx")
    with pytest.raises(ValueError, match="商品"):
        sales_pricebook.import_quote_template(str(home / "t.xlsx"))


def test_system_entries_do_not_hide_waiting_customer(home):
    pricebook(home)
    make()
    c(action="log", customer="南岭建设", date="2026-10-08", channel="微信", who="客户", kind="问题", text="能便宜点吗?")
    c(action="quote", customer="南岭建设", items=[{"name": "实施服务", "qty": 1}])
    c(action="upsert", name="南岭建设", stage="已报价", confirmed=True)
    rules = [x["rule"] for x in customer_radar.scan(date(2026, 10, 12))["cards"]]
    assert "reply" in rules  # 改阶段 / 生成报价单不等于回过客户
    assert c(action="log", customer="南岭建设", channel="系统", who="我方", kind="信息", text="x")["ok"] is False


# ── 销售周报 ──

def test_weekly_report_counts(home):
    pricebook(home)
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    make("南岭建设", deal_size="12万/年")
    make("恒远检测", stage="商务谈判", deal_size="¥80,000", contacts=[{"name": "陈立", "email": "chen@hengyuan-test.cn"}])
    c(action="upsert", name="南岭建设", stage="已报价", confirmed=True)
    c(action="log", customer="南岭建设", date=monday.isoformat(), channel="邮件", who="客户", kind="问题", text="报价含不含税?")
    c(action="log", customer="南岭建设", date=monday.isoformat(), channel="邮件", who="我方", kind="报价", text="v1 合计 ¥12万")
    c(action="log", customer="恒远检测", date=(monday - timedelta(days=3)).isoformat(), channel="微信", who="客户", kind="信息", text="上周的事")
    c(action="quote", customer="南岭建设", items=[{"name": "鲶鱼企业版", "spec": "50 人", "qty": 1}])

    r = c(action="weekly")
    assert r["ok"] and r["customers"] == 2 and sorted(r["new_customers"]) == ["南岭建设", "恒远检测"]
    funnel = {f["stage"]: f for f in r["funnel"]}
    assert funnel["已报价"]["count"] == 1 and funnel["已报价"]["deal_total"] == 120000
    assert funnel["商务谈判"]["deal_total"] == 80000
    assert r["stage_changes"] == [{"customer": "南岭建设", "from": "需求确认", "to": "已报价"}]
    assert [x["version"] for x in r["quotes_made"]] == [1] and len(r["quotes_sent"]) == 1
    nanling = next(x for x in r["active_customers"] if x["customer"] == "南岭建设")
    assert nanling["from_customer"] == 1 and nanling["from_us"] == 1
    wb = load_workbook(r["file"])
    assert wb.sheetnames == ["概览", "客户明细", "待跟进"]
    assert Path(r["file"]).name == f"销售周报-{monday.isoformat()}.xlsx"


def test_weekly_without_customers(home):
    r = c(action="weekly")
    assert not r["ok"] and "客户档案" in r["error"]


# ── Companion 设置页走的 RPC ──

def rpc(method, params=None):
    from catfish_tool_bridge import server

    return asyncio.run(server._handle_request({"id": 1, "method": method, "params": params or {}}))


def test_rpc_pricebook_and_quote_settings(home):
    p = xlsx(home / "p.xlsx", [["商品", "标价"], ["鲶鱼个人版", 999]])
    assert rpc("pim/sales_pricebook_import", {"path": str(p)})["result"]["count"] == 1
    assert rpc("pim/sales_pricebook_get")["result"]["items"][0]["name"] == "鲶鱼个人版"
    assert "error" in rpc("pim/sales_pricebook_import", {"path": str(home / "nope.xlsx")})
    s = rpc("pim/sales_quote_settings_set", {"settings": {"company": "示例科技有限公司", "valid_days": 15}})["result"]
    assert s["company"] == "示例科技有限公司" and s["valid_days"] == 15 and "含税" in s["terms"]
    assert rpc("pim/sales_quote_settings_get")["result"]["custom_template"] is False
    assert "error" in rpc("pim/sales_quote_settings_set", {"settings": {"valid_days": "x"}})


# ── 阶段 3: 发布给主管 (sales_publish) ──

def test_publish_summary_has_no_chat_text(home, monkeypatch):
    from catfish_tool_bridge import sales_publish

    sent = []

    def fake_http(method, url, payload):
        sent.append((method, url, payload))
        if url.endswith("/unpublish"):
            return 200, {"ok": True}
        return 200, {"file_id": "f1", "namespace": "dept/销售部"}

    monkeypatch.setattr(sales_publish, "_http", fake_http)
    with pytest.raises(ValueError, match="没有客户"):
        sales_publish.publish()
    make(deal_size="12万/年")
    c(action="log", customer="南岭建设", date="2026-10-08", channel="微信", who="客户", kind="问题", text="私有化部署要几台服务器")
    pv = sales_publish.preview()["summary"]
    assert pv["customers"][0]["flags"] == ["客户在等你回复"] and pv["customers"][0]["deal_value"] == 120000

    r = sales_publish.publish()
    assert r["ok"] and r["file_id"] == "f1"
    method, url, payload = sent[-1]
    assert method == "POST" and url.endswith("/v1/wiki/documents") and "file_id" not in payload
    assert payload["filename"] == "catfish/sales-summary.md"  # hub 靠这个认, 对同事隐藏
    blob = json.dumps(payload, ensure_ascii=False)
    assert "南岭建设" in blob and "私有化部署" not in blob and "zhou@nanling-build.cn" not in blob  # 不带原话 / 联系方式
    line = next(x for x in payload["frontmatter_yaml"].splitlines() if x.startswith("sales_summary:"))
    assert json.loads(line[len("sales_summary:"):])["funnel"][1]["stage"] == "需求确认"

    sales_publish.publish()
    assert sent[-1][2]["file_id"] == "f1"  # 重发覆盖同一条
    assert sales_publish.unpublish()["unpublished"] and sent[-1][1].endswith("/f1/unpublish")
    assert sales_publish.load_state() == {}
