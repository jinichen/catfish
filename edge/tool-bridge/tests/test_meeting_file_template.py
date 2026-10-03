"""公司固定的 Word / Excel 纪要模版: 认空 + 填写 (10/3)。

模版都在测试里现造, 照常见的单位纪要表的样子: 标签格 + 空格、合并大格、表头 + 空行
的待办表、待办表下面还有「记录人」。重点验: 认对、填对位置、格式和其它内容不动、
待办比空行多时加行且下面的格子不错位。
"""
from __future__ import annotations

from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")

from catfish_tool_bridge import meeting_file_template as ft  # noqa: E402

TODOS = [["1", "交二期材料", "张三", "2026-10-10"], ["2", "约评审会", "李四", "2026-10-12"],
         ["3", "整理报价", "王五", ""], ["4", "通知法务", "赵六", "2026-10-15"]]


def _form_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("XX 公司会议纪要", 1)
    t = d.add_table(rows=4, cols=4)
    t.style = "Table Grid"
    for (r, c), text in {(0, 0): "会议名称", (0, 2): "会议时间：", (1, 0): "主持人", (1, 2): "地点",
                         (2, 0): "参会人员"}.items():
        t.cell(r, c).text = text
    t.cell(2, 1).merge(t.cell(2, 3))
    t.cell(3, 0).merge(t.cell(3, 3)).text = "会议内容"
    big = d.add_table(rows=1, cols=1)  # 「会议内容」下面那一大格 —— 这里用单独一行表模拟
    big.style = "Table Grid"
    d.add_paragraph("三、下一步工作")
    todo = d.add_table(rows=4, cols=4)
    todo.style = "Table Grid"
    for i, h in enumerate(["序号", "事项", "负责人", "完成时限"]):
        todo.cell(0, i).text = h
    tail = d.add_table(rows=1, cols=2)
    tail.cell(0, 0).text = "记录人"
    d.add_paragraph("备注：本纪要发至各部门。")
    d.save(str(path))
    return path


def test_docx_form_slots_detected(tmp_path):
    slots = ft.inspect(_form_docx(tmp_path / "t.docx"))
    labels = {s["label"]: s for s in slots}
    for name in ("会议名称", "会议时间", "主持人", "地点", "参会人员", "记录人"):
        assert labels[name]["kind"] == "text", name
    todo = labels["三、下一步工作"]
    assert todo["kind"] == "table" and todo["columns"] == ["序号", "事项", "负责人", "完成时限"]
    assert "备注" not in labels, "「备注：本纪要发至各部门。」冒号后面有字, 不是空"


def test_docx_form_filled_in_place_and_rows_added(tmp_path):
    src = _form_docx(tmp_path / "t.docx")
    slots = ft.inspect(src)
    by = {s["label"]: s["id"] for s in slots}
    values = {by["会议名称"]: "资质集采周会", by["会议时间"]: "2026-10-01", by["参会人员"]: "陈鸿波、林达华",
              by["三、下一步工作"]: TODOS, by["记录人"]: "小鲶"}
    out = tmp_path / "out.docx"
    ft.fill(src, out, slots, values)
    d = docx.Document(str(out))
    t = d.tables[0]
    assert t.cell(0, 1).text == "资质集采周会" and t.cell(0, 3).text == "2026-10-01"
    assert t.cell(2, 1).text == "陈鸿波、林达华"
    todo = d.tables[2]
    assert len(todo.rows) == 5, "3 个空行不够 4 个待办, 照样子加一行"
    assert [c.text for c in todo.rows[4].cells] == TODOS[3]
    assert todo.style.name == "Table Grid"
    assert d.tables[3].cell(0, 1).text == "小鲶", "待办表下面的「记录人」没被加行挤错位"
    assert d.paragraphs[0].text == "XX 公司会议纪要" and d.paragraphs[-1].text == "备注：本纪要发至各部门。"


def test_docx_colon_paragraphs(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("参会人员：").bold = True
    p.add_run("________")
    d.add_paragraph("会议决定:")
    src = tmp_path / "c.docx"
    d.save(str(src))
    slots = ft.inspect(src)
    assert [s["label"] for s in slots] == ["参会人员", "会议决定"]
    out = tmp_path / "c-out.docx"
    ft.fill(src, out, slots, {slots[0]["id"]: "张三、李四", slots[1]["id"]: ["一、同意二期方案", "二、周五前交材料"]})
    paras = docx.Document(str(out)).paragraphs
    assert paras[0].text == "参会人员：张三、李四", "下划线占位去掉"
    assert paras[0].runs[-1].bold, "填的字跟标签同样式"
    assert paras[1].text == "会议决定:一、同意二期方案\n二、周五前交材料"


def test_docx_placeholders_and_repeating_row(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("会议：")
    p.add_run("{{会议")  # Word 常把占位符拆进好几个 run
    p.add_run("名称}}")
    t = d.add_table(rows=2, cols=3)
    for i, h in enumerate(["事项", "负责人", "时限"]):
        t.cell(0, i).text = h
    for i, h in enumerate(["{{事项}}", "{{负责人}}", "{{时限}}"]):
        t.cell(1, i).text = h
    d.add_paragraph("参会：{{参会人员}}")
    src = tmp_path / "p.docx"
    d.save(str(src))
    slots = ft.inspect(src)
    kinds = {s["label"]: s["kind"] for s in slots}
    assert kinds["会议名称"] == "text" and kinds["参会人员"] == "text"
    table = next(s for s in slots if s["kind"] == "table")
    assert table["columns"] == ["事项", "负责人", "时限"]
    by = {s["label"]: s["id"] for s in slots}
    out = tmp_path / "p-out.docx"
    ft.fill(src, out, slots, {by["会议名称"]: "周会", by["参会人员"]: "张三", table["id"]: [r[1:] for r in TODOS[:2]]})
    d2 = docx.Document(str(out))
    assert d2.paragraphs[0].text == "会议：周会" and d2.paragraphs[1].text == "参会：张三"
    assert [[c.text for c in r.cells] for r in d2.tables[0].rows] == [
        ["事项", "负责人", "时限"], ["交二期材料", "张三", "2026-10-10"], ["约评审会", "李四", "2026-10-12"]]


def _form_xlsx(path: Path) -> Path:
    from openpyxl.styles import Border, Font, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "纪要"
    ws["A1"] = "XX 公司会议纪要"
    ws.merge_cells("A1:D1")
    ws["A2"], ws["C2"] = "会议名称", "会议时间"
    ws["A3"] = "参会人员"
    ws.merge_cells("B3:D3")
    ws["A4"] = "会议内容"
    ws.merge_cells("A5:D5")  # 下面一整行大格
    ws["A6"] = "下一步工作"
    thin = Side(style="thin")
    for i, h in enumerate(["序号", "事项", "负责人", "完成时限"], 1):
        c = ws.cell(7, i, h)
        c.font = Font(bold=True)
    for r in (8, 9):
        for col in range(1, 5):
            ws.cell(r, col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws["A11"] = "记录人"
    ws.merge_cells("B11:C11")
    wb.save(str(path))
    return path


def test_xlsx_form_detect_and_fill(tmp_path):
    src = _form_xlsx(tmp_path / "t.xlsx")
    slots = ft.inspect(src)
    by = {s["label"]: s for s in slots}
    assert {"会议名称", "会议时间", "参会人员", "会议内容", "记录人"} <= set(by)
    todo = by["下一步工作"]
    assert todo["columns"] == ["序号", "事项", "负责人", "完成时限"]
    assert "XX 公司会议纪要" not in by, "合并的大标题不是要填的空"
    values = {by["会议名称"]["id"]: "资质集采周会", by["会议时间"]["id"]: "2026-10-01",
              by["参会人员"]["id"]: "陈鸿波、林达华", by["会议内容"]["id"]: "一、过二期\n二、定分工",
              todo["id"]: TODOS, by["记录人"]["id"]: "小鲶"}
    out = tmp_path / "out.xlsx"
    ft.fill(src, out, slots, values)
    ws = openpyxl.load_workbook(str(out))["纪要"]
    assert ws["B2"].value == "资质集采周会" and ws["D2"].value == "2026-10-01"
    assert ws["B3"].value == "陈鸿波、林达华"
    assert ws["A5"].value == "一、过二期\n二、定分工" and ws["A5"].alignment.wrap_text
    assert [ws.cell(11, c).value for c in range(1, 5)] == TODOS[3], "两个空行 + 一行空白不够, 插了行"
    assert ws.cell(11, 2).border.top.style == "thin", "新插的行照原来空行的样子"
    assert ws["A13"].value == "记录人" and ws["B13"].value == "小鲶", "下面的「记录人」跟着下移, 没写错地方"
    assert "B13:C13" in {str(r) for r in ws.merged_cells.ranges}, "合并格跟着挪"
    assert ws["A1"].value == "XX 公司会议纪要"


def test_xlsx_placeholders_repeating_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "会议: {{会议名称}}"
    ws.append(["事项", "负责人"])
    ws.append(["{{事项}}", "{{负责人}}"])
    ws["A5"] = "记录: {{记录人}}"
    src = tmp_path / "p.xlsx"
    wb.save(str(src))
    slots = ft.inspect(src)
    by = {s["label"]: s["id"] for s in slots}
    table = next(s for s in slots if s["kind"] == "table")
    out = tmp_path / "p-out.xlsx"
    ft.fill(src, out, slots, {by["会议名称"]: "周会", by["记录人"]: "小鲶",
                              table["id"]: [["交材料", "张三"], ["约会", "李四"], ["报价", "王五"]]})
    ws2 = openpyxl.load_workbook(str(out)).active
    assert ws2["A1"].value == "会议: 周会"
    assert [[ws2.cell(r, c).value for c in (1, 2)] for r in (3, 4, 5)] == [["交材料", "张三"], ["约会", "李四"], ["报价", "王五"]]
    assert ws2["A7"].value == "记录: 小鲶"


def test_rejects_old_formats_and_empty_templates(tmp_path):
    with pytest.raises(ft.TemplateError, match="另存为"):
        ft.inspect(tmp_path / "x.doc")
    d = docx.Document()
    d.add_paragraph("一段说明文字, 没有要填的地方。")
    d.save(str(tmp_path / "e.docx"))
    with pytest.raises(ft.TemplateError, match="占位符"):
        ft.inspect(tmp_path / "e.docx")


def test_summarize_with_word_template_end_to_end(tmp_path, monkeypatch):
    """生成纪要选 Word 模版: 第二次调大模型拿到"要填的空", 回的值按 id 填进新文件, 原模版不动。"""
    import asyncio
    import json

    from catfish_tool_bridge import meeting_minutes as mm
    from tests.test_meeting_minutes import REPLY, SEGS, _fake_llm, _meeting

    src = _form_docx(tmp_path / "公司模版.docx")
    before = src.read_bytes()
    slots = ft.inspect(src)
    by = {s["label"]: s["id"] for s in slots}
    fill_reply = json.dumps({by["会议名称"]: "资质集采周会", by["三、下一步工作"]: [["1", "交材料", "张三", "2026-10-10"]],
                             "s999": "不认识的 id 丢掉"}, ensure_ascii=False)
    d = _meeting(tmp_path, SEGS, {"0": "陈鸿波"})
    calls = _fake_llm(monkeypatch, [REPLY, fill_reply])
    doc = asyncio.run(mm.handle_summarize({
        "catfish_home": str(tmp_path), "gateway_url": "http://gw.test", "auth_token": "t",
        "meeting_id": "mtg_20261001-100000_ab12",
        "template": {"id": "tpl_w", "name": "公司纪要表", "kind": "docx", "path": str(src)},
    }))
    sent = calls[1][1]["content"]
    assert "会议时间" in sent and "完成时限" in sent and '"会议名称": "资质集采周会"' in sent
    assert '"loc"' not in sent, "文件里的定位信息不发给大模型"
    out = Path(doc["output_file"])
    assert out.name == "资质集采周会 会议纪要.docx" and out.parent == d
    filled = docx.Document(str(out))
    assert filled.tables[0].cell(0, 1).text == "资质集采周会"
    assert [c.text for c in filled.tables[2].rows[1].cells] == ["1", "交材料", "张三", "2026-10-10"]
    assert src.read_bytes() == before, "原模版文件不能被改"
    assert doc["template"] == {"id": "tpl_w", "name": "公司纪要表", "kind": "docx"}
    assert (d / "minutes.md").read_text(encoding="utf-8").startswith("# 资质集采周会"), "界面 / 知识库照常用缺省版"


def test_missing_template_file_and_inspect_rpc(tmp_path):
    from catfish_tool_bridge import meeting_minutes as mm
    from catfish_tool_bridge import meeting_minutes_template as mt

    with pytest.raises(ValueError, match="找不到"):
        mt.validate_template({"kind": "docx", "path": str(tmp_path / "gone.docx")})
    src = _form_xlsx(tmp_path / "t.xlsx")
    got = mm.handle_template_inspect({"path": str(src)})["slots"]
    assert any(s["kind"] == "table" for s in got) and all("loc" not in s for s in got)


def test_todo_tables_filled_from_structured_items_not_model():
    cols = ["序号", "工作事项", "责任人", "完成时限", "备注"]
    assert ft.is_todo_table(cols) and not ft.is_todo_table(["议题", "发言人"])
    items = [{"task": "交材料", "owner": "张三", "due": "2026-10-10"}, {"task": "约会", "owner": "", "due": ""}]
    assert ft.todo_rows(cols, items) == [["1", "交材料", "张三", "2026-10-10", ""], ["2", "约会", "", "", ""]]
    slots = [{"id": "s1", "kind": "table", "columns": cols}, {"id": "s2", "kind": "text", "label": "x"}]
    got = ft.apply_todos(slots, {"s1": [], "s2": "v"}, items)
    assert got["s1"][0][1] == "交材料" and got["s2"] == "v"
    assert ft.apply_todos(slots, {"s1": [["模型给的"]]}, [])["s1"] == [["模型给的"]], "没抽到待办就用模型的"


def test_docx_content_box_as_separate_table(tmp_path):
    src = _form_docx(tmp_path / "t.docx")  # 「会议内容」在表 0 最后一行, 大空格是紧跟的单格表
    slots = ft.inspect(src)
    content = next(s for s in slots if s["label"] == "会议内容")
    out = tmp_path / "o.docx"
    ft.fill(src, out, slots, {content["id"]: "一、过二期\n二、定分工"})
    box = docx.Document(str(out)).tables[1].cell(0, 0)
    assert [p.text for p in box.paragraphs] == ["一、过二期", "二、定分工"]
