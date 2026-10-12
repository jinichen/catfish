"""销售周报 (销售跟进阶段 2, docs/SALES-FOLLOWUP-PLAN.md §10.3)。

catfish_customer action=weekly: 只读客户档案 + 报价台账, 确定性统计本周 (周一到 date), 出
~/.catfish/outputs/<日期>/销售周报-<周一>.xlsx (概览 / 客户明细 / 待跟进 三页) + 给聊天用的摘要。
数字全是数出来的, 不让模型估; 模型只负责把摘要念给员工、按员工口吻补几句。
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from . import advisor_io, customer_radar, customers, sales_quote

_STAGE_CHANGE = re.compile(r"^阶段 (.+?) → (.+)$")


def deal_value(s: Any) -> float | None:
    """「12万/年」「¥80,000」→ 数; 估不出就不计入合计。"""
    m = re.search(r"(\d+(?:\.\d+)?)\s*(万)?", str(s or "").replace(",", ""))
    if not m:
        return None
    return float(m.group(1)) * (10000 if m.group(2) else 1)


def _in(day: Any, start: date, end: date) -> bool:
    try:
        d = date.fromisoformat(str(day)[:10])
    except ValueError:
        return False
    return start <= d <= end


def build(today: date) -> dict[str, Any]:
    start = today - timedelta(days=today.weekday())
    tpl = customers.load_template()
    all_c = customers.all_customers()
    ledger = sales_quote.load_ledger()

    funnel = []
    for s in tpl["stages"]:
        rows = [c for c in all_c if c.get("stage") == s["name"]]
        vals = [v for v in (deal_value(c.get("deal_size")) for c in rows) if v is not None]
        funnel.append({"stage": s["name"], "count": len(rows), "deal_total": sum(vals)})

    per_customer, stage_changes, quotes_sent = [], [], []
    for c in all_c:
        week = [e for e in c["timeline"] if _in(e["date"], start, today)]
        talk = [e for e in week if e["channel"] != customers.SYSTEM_CHANNEL]
        changes = [m.groups() for e in week if e["channel"] == customers.SYSTEM_CHANNEL
                   for m in [_STAGE_CHANGE.match(e["text"])] if m]
        for frm, to in reversed(changes):  # 时间线新的在上, 周报按发生顺序
            stage_changes.append({"customer": c["name"], "from": frm, "to": to})
        for e in talk:
            if e["who"] == "我方" and e["kind"] == "报价":
                quotes_sent.append({"customer": c["name"], "date": e["date"], "text": e["text"]})
        new = _in(c.get("created"), start, today)
        if talk or changes or new:
            per_customer.append({
                "customer": c["name"], "stage": c.get("stage"), "owner": c.get("owner") or "",
                "new": new, "from_customer": sum(1 for e in talk if e["who"] == "客户"),
                "from_us": sum(1 for e in talk if e["who"] == "我方"),
                "stage_change": " / ".join(f"{a}→{b}" for a, b in reversed(changes)),
                "highlights": [f"{e['who']}{e['kind']}: {e['text']}" for e in talk[:3]],
                "next_action": c.get("next_action") or "", "next_action_due": c.get("next_action_due") or "",
            })

    quotes_made = [x for x in ledger if _in(x.get("date"), start, today)]
    radar = customer_radar.scan(today)["cards"]
    active = {s["name"] for s in tpl["stages"] if s["cold_days"] > 0}
    next_week = sorted(
        ({"customer": c["name"], "stage": c.get("stage"), "next_action": c["next_action"], "due": c.get("next_action_due") or ""}
         for c in all_c if c.get("next_action") and c.get("stage") in active),
        key=lambda x: x["due"] or "9999")
    return {
        "week_start": start.isoformat(), "week_end": today.isoformat(), "template": tpl["name"],
        "customers": len(all_c),
        "new_customers": [x["customer"] for x in per_customer if x["new"]],
        "funnel": funnel, "stage_changes": stage_changes,
        "quotes_made": [{k: x.get(k) for k in ("customer", "version", "date", "total")} for x in quotes_made],
        "quotes_sent": quotes_sent,
        "followup": dict(Counter(x["rule"] for x in radar)), "followup_cards": radar,
        "per_customer": per_customer, "next_week": next_week,
    }


def _write_xlsx(path: Path, r: dict[str, Any]) -> None:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415

    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDEBF7")

    def table(ws: Any, row: int, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> int:
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.font, cell.fill, cell.border = Font(bold=True), head, box
        for vals in rows:
            row += 1
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row, c, v)
                cell.border = box
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[chr(64 + i)].width = max(ws.column_dimensions[chr(64 + i)].width or 0, w)
        sales_quote._fit_width(ws)
        return row + 3  # 空一行, 下一节标题写在 row+2

    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    ws["A1"] = f"销售周报 {r['week_start']} ~ {r['week_end']}"
    ws["A1"].font = Font(size=14, bold=True)
    sent_n, made_n = len(r["quotes_sent"]), len(r["quotes_made"])
    fu = r["followup"]
    row = table(ws, 3, ["指标", "本周"], [
        ["客户总数", r["customers"]], ["本周新建客户", len(r["new_customers"])],
        ["阶段变化", len(r["stage_changes"])], ["生成报价单", made_n], ["发出报价", sent_n],
        ["客户在等回复", fu.get("reply", 0)], ["答应的事到期", fu.get("promise", 0)],
        ["报价没回音", fu.get("quote", 0)], ["客户冷了", fu.get("cold", 0)],
    ], [18, 12, 14])
    ws.cell(row - 1, 1, "销售漏斗").font = Font(bold=True)
    row = table(ws, row, ["阶段", "客户数", "预计金额"],
                [[f["stage"], f["count"], f["deal_total"] or ""] for f in r["funnel"]])
    for cells in ws.iter_rows(min_col=3, max_col=3):
        for cell in cells:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    if r["stage_changes"]:
        ws.cell(row - 1, 1, "本周阶段变化").font = Font(bold=True)
        table(ws, row, ["客户", "从", "到"], [[x["customer"], x["from"], x["to"]] for x in r["stage_changes"]])

    ws2 = wb.create_sheet("客户明细")
    table(ws2, 1, ["客户", "阶段", "负责人", "新建", "客户来信", "我方发出", "阶段变化", "本周要点", "下一步", "截止"],
          [[x["customer"], x["stage"], x["owner"], "是" if x["new"] else "", x["from_customer"], x["from_us"],
            x["stage_change"], "\n".join(x["highlights"]), x["next_action"], x["next_action_due"]] for x in r["per_customer"]],
          [16, 10, 10, 6, 9, 9, 16, 48, 28, 12])

    ws3 = wb.create_sheet("待跟进")
    table(ws3, 1, ["客户", "提醒", "依据", "已过天数", "下一步"],
          [[x["customer"], x["label"], x["quote"], x["age_days"], x.get("next_action") or ""] for x in r["followup_cards"]],
          [16, 18, 48, 10, 28])
    wb.save(path)


def action_weekly(args: dict[str, Any]) -> dict[str, Any]:
    raw = str(args.get("date") or "").strip()
    try:
        today = date.fromisoformat(raw[:10]) if raw else date.today()
    except ValueError:
        raise ValueError(f"日期格式不对 (要 YYYY-MM-DD): {raw}") from None
    r = build(today)
    if not r["customers"]:
        return {"ok": False, "error": "还没有客户档案, 出不了销售周报。先让员工把在跟的客户建起来 (action=upsert)"}
    path = advisor_io.today_outputs_dir() / f"销售周报-{r['week_start']}.xlsx"
    _write_xlsx(path, r)
    summary = {k: r[k] for k in ("week_start", "week_end", "customers", "new_customers", "funnel", "stage_changes",
                                 "quotes_made", "quotes_sent", "followup", "next_week")}
    summary["active_customers"] = [{k: x[k] for k in ("customer", "stage", "from_customer", "from_us", "highlights")}
                                   for x in r["per_customer"][:15]]
    return {"ok": True, "file": str(path), **summary,
            "next": "把文件路径给员工, 再用三五句话讲本周: 漏斗变化、发了哪些报价、哪几个客户要优先跟。数字照上面念, 不要自己估。"}
