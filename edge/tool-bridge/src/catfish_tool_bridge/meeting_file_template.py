"""公司固定的 Word / Excel 纪要模版: 认出要填的空, 填好另存 (10/3)。

很多单位的纪要是定死的 .docx / .xlsx 表单 (表头、边框、合并单元格、落款都固定),
员工要的是"把这张表填好", 不是另起一份 Markdown。这里只做两件事, 都不碰原格式:

1. inspect(path) → slots: 文件里哪些地方要填。两种认法:
   - 占位符 (员工自己在模版里写的, 优先): 单元格 / 段落里的 {{会议名称}};
     表格的一行里有两个以上占位符 = "每个待办一行"的循环行。
   - 自动认 (模版一个字没改也能用):
       · 表格: 标签格 (「会议时间」「主持人：」) 右边 (没有就下边) 是空格
       · 表格: 一行表头 (≥2 格有字) 下面跟着整行空行 → 多行表 (待办 / 议题清单)
       · 段落: 以冒号结尾的「参会人员：」, 填在冒号后面
       · Excel 同理, 合并单元格按左上角那格算
2. fill(src, dst, slots, values): 按 slots 把值写进去另存。写进原来那个格子 / 段落, 字体
   边框照旧; 待办比空行多就照最后一行的样子加行。

大模型只负责"每个空填什么" (meeting_minutes.py 里那一次调用), 落到文件上全是确定的代码。
"""
from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_LABEL_MAX = 20
_TRAILING_BLANK_RE = re.compile(r"[\s_＿]+$")
SUPPORTED = (".docx", ".xlsx")


class TemplateError(ValueError):
    pass


def _clean_label(text: str) -> str:
    return re.sub(r"[：:\s]+$", "", text.strip())


def _is_label(text: str) -> bool:
    t = text.strip()
    return 0 < len(t) <= _LABEL_MAX and not re.fullmatch(r"[\d\s.、,，-]+", t) and not PLACEHOLDER_RE.search(t)


def inspect(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise TemplateError(f"只支持 .docx / .xlsx; {ext or '这个文件'} 请先用 Word / Excel / WPS 另存为 .docx 或 .xlsx")
    slots = _inspect_docx(path) if ext == ".docx" else _inspect_xlsx(path)
    if not slots:
        raise TemplateError(
            "没在模版里找到要填的地方。可以在要填的位置写上占位符, 比如 {{会议名称}} {{参会人员}}; "
            "待办表在一行里写 {{事项}} {{负责人}} {{完成时限}}, 每个待办会复制出一行。"
        )
    for i, s in enumerate(slots, 1):
        s["id"] = f"s{i}"
    return slots


# ── Word ──────────────────────────────────────────────────────────────


def _docx_paragraph_text(p) -> str:
    return "".join(r.text for r in p.runs)


def _unique_cells(row) -> list:
    """一行里的格子, 合并单元格只算一次 (python-docx 对横向合并会重复返回同一个格子)。"""
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


def _inspect_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document  # noqa: PLC0415 —— hermes venv 里有 (hermes-extra-packages)

    doc = Document(str(path))
    has_placeholders = any(PLACEHOLDER_RE.search(_docx_paragraph_text(p)) for p in doc.paragraphs) or any(
        PLACEHOLDER_RE.search(c.text) for t in doc.tables for row in t.rows for c in row.cells
    )
    slots: list[dict[str, Any]] = []
    used_paras: set = set()
    if has_placeholders:
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                names = [n for c in _unique_cells(row) for n in PLACEHOLDER_RE.findall(c.text)]
                if len(names) >= 2:
                    slots.append({"kind": "table", "label": _table_label(doc, t, ti), "columns": names,
                                  "loc": {"type": "docx_row_template", "table": ti, "row": ri}})
                else:
                    slots += [{"kind": "text", "label": n, "loc": {"type": "placeholder", "name": n}} for n in names]
        for p in doc.paragraphs:
            slots += [{"kind": "text", "label": n, "loc": {"type": "placeholder", "name": n}}
                      for n in PLACEHOLDER_RE.findall(_docx_paragraph_text(p))]
        return _dedupe_placeholders(slots)

    for ti, t in enumerate(doc.tables):
        rows = [_unique_cells(r) for r in t.rows]
        used: set[tuple[int, int]] = set()
        # 1) 表头 + 空行 → 多行表
        for ri, cells in enumerate(rows):
            texts = [c.text.strip() for c in cells]
            if len(cells) >= 2 and all(texts) and all(_is_label(x) for x in texts):
                empties = []
                for rj in range(ri + 1, len(rows)):
                    if len(rows[rj]) == len(cells) and not any(c.text.strip() for c in rows[rj]):
                        empties.append(rj)
                    else:
                        break
                if empties:
                    slots.append({"kind": "table", "label": _table_label(doc, t, ti, texts, used_paras), "columns": texts,
                                  "loc": {"type": "docx_rows", "table": ti, "header": ri, "rows": empties}})
                    used |= {(r, c) for r in [ri, *empties] for c in range(len(cells))}
        # 2) 标签格 + 右边 (或下边) 的空格
        for ri, cells in enumerate(rows):
            for ci, c in enumerate(cells):
                if (ri, ci) in used or not _is_label(c.text):
                    continue
                if ci + 1 < len(cells) and (ri, ci + 1) not in used and not cells[ci + 1].text.strip():
                    target = (ri, ci + 1)
                elif (ci + 1 >= len(cells) and ri + 1 < len(rows) and len(rows[ri + 1]) == 1
                      and not rows[ri + 1][0].text.strip()):
                    target = (ri + 1, 0)  # 「会议内容」下面一整行的大空格
                elif len(cells) == 1 and ri == len(rows) - 1 and _next_empty_box(doc, t, ti) is not None:
                    # 「会议内容」是表格最后一行, 大空格是紧跟着的另一张单格表 (Word 里常这么画)
                    nti = _next_empty_box(doc, t, ti)
                    slots.append({"kind": "text", "label": _clean_label(c.text),
                                  "loc": {"type": "docx_cell", "table": nti, "row": 0, "col": 0}})
                    used.add((ri, ci))
                    continue
                else:
                    continue
                if target in used:
                    continue
                used |= {(ri, ci), target}
                slots.append({"kind": "text", "label": _clean_label(c.text),
                              "loc": {"type": "docx_cell", "table": ti, "row": target[0], "col": target[1]}})
    # 3) 段落「参会人员：」
    for pi, p in enumerate(doc.paragraphs):
        if p._p in used_paras:
            continue
        text = _docx_paragraph_text(p)
        stripped = _TRAILING_BLANK_RE.sub("", text)
        if stripped.endswith(("：", ":")) and _is_label(stripped):
            slots.append({"kind": "text", "label": _clean_label(stripped), "loc": {"type": "docx_after_colon", "para": pi}})
    return slots


def _next_empty_box(doc, table, ti: int) -> int | None:
    """紧跟在 table 后面 (中间没有段落) 的那张表如果只有一个空格子, 返回它的下标。"""
    nxt = table._tbl.getnext()
    if nxt is None or not nxt.tag.endswith("}tbl") or ti + 1 >= len(doc.tables):
        return None
    box = doc.tables[ti + 1]
    if box._tbl is not nxt or len(box.rows) != 1 or len(_unique_cells(box.rows[0])) != 1:
        return None
    return ti + 1 if not box.rows[0].cells[0].text.strip() else None


def _table_label(doc, table, ti: int, header: list[str] | None = None, used_paras: set | None = None) -> str:
    """多行表叫什么: 表格前一段的文字 (「三、下一步工作」), 没有就用表头拼。
    用作表名的那段记进 used_paras, 后面"冒号结尾的段落"不再把它当成一个空。"""
    prev = table._tbl.getprevious()
    while prev is not None and prev.tag.endswith("}p"):
        text = "".join(t.text or "" for t in prev.iter() if t.tag.endswith("}t")).strip()
        if text:
            if used_paras is not None:
                used_paras.add(prev)
            return _clean_label(text)[:_LABEL_MAX]
        prev = prev.getprevious()
    return "/".join(header or [])[:_LABEL_MAX] or f"表格{ti + 1}"


def _dedupe_placeholders(slots: list[dict]) -> list[dict]:
    seen, out = set(), []
    for s in slots:
        key = s["label"] if s["loc"]["type"] == "placeholder" else id(s)
        if key not in seen:
            seen.add(key)
            out.append(s)
    return out


def _set_cell_text(cell, text: str) -> None:
    lines = str(text).split("\n")
    p0 = cell.paragraphs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    if p0.runs:
        p0.runs[0].text = lines[0]
    else:
        p0.add_run(lines[0])
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for line in lines[1:]:
        cell.add_paragraph(line, style=p0.style)


def _replace_in_paragraph(p, values: dict[str, str]) -> None:
    full = _docx_paragraph_text(p)
    if not PLACEHOLDER_RE.search(full) or not p.runs:
        return
    new = PLACEHOLDER_RE.sub(lambda m: str(values.get(m.group(1), "")), full)
    # 占位符常被 Word 拆在好几个 run 里: 整段文字放进第一个 run (保它的字体), 其余清空
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def _fill_docx(src: Path, dst: Path, slots: list[dict], values: dict[str, Any]) -> None:
    from docx import Document  # noqa: PLC0415

    doc = Document(str(src))
    ph_values = {s["label"]: _as_text(values.get(s["id"], "")) for s in slots if s["loc"]["type"] == "placeholder"}
    # 表格里的写入从下往上做: 某张表加了行, 它下面的格子行号会变, 先写下面的就不受影响
    def _pos(s: dict) -> tuple[int, int]:
        loc = s["loc"]
        return (loc.get("table", -1), loc.get("row", (loc.get("rows") or [-1])[0]))

    for s in sorted(slots, key=_pos, reverse=True):
        loc = s["loc"]
        if loc["type"] == "docx_row_template":
            table = doc.tables[loc["table"]]
            tmpl = table.rows[loc["row"]]._tr
            for row in _as_rows(values.get(s["id"]), len(s["columns"])):
                new_tr = copy.deepcopy(tmpl)
                tmpl.addprevious(new_tr)
                from docx.table import _Row  # noqa: PLC0415
                r = _Row(new_tr, table)
                for c in _unique_cells(r):
                    for p in c.paragraphs:
                        _replace_in_paragraph(p, dict(zip(s["columns"], row)))
            tmpl.getparent().remove(tmpl)
        elif loc["type"] == "docx_rows":
            table = doc.tables[loc["table"]]
            rows = _as_rows(values.get(s["id"]), len(s["columns"]))
            targets = list(loc["rows"])
            last = table.rows[targets[-1]]._tr
            for _ in range(max(0, len(rows) - len(targets))):  # 空行不够: 照最后一个空行的样子加
                new_tr = copy.deepcopy(last)
                last.addnext(new_tr)
                last = new_tr
            all_rows = list(table.rows)
            start = targets[0]
            for i, row in enumerate(rows):
                for c, text in zip(_unique_cells(all_rows[start + i]), row):
                    _set_cell_text(c, text)
        elif loc["type"] == "docx_cell":
            cell = _unique_cells(doc.tables[loc["table"]].rows[loc["row"]])[loc["col"]]
            _set_cell_text(cell, _as_text(values.get(s["id"], "")))
        elif loc["type"] == "docx_after_colon":
            p = doc.paragraphs[loc["para"]]
            for r in reversed(p.runs):  # 去掉冒号后面的下划线 / 空格占位
                stripped = _TRAILING_BLANK_RE.sub("", r.text)
                if stripped != r.text or not r.text:
                    r.text = stripped
                    if stripped:
                        break
                else:
                    break
            label_run = next((r for r in reversed(p.runs) if r.text), None)  # 去掉下划线后最后一个有字的 = 标签
            run = p.add_run(_as_text(values.get(s["id"], "")))
            if label_run is not None and label_run._r.rPr is not None:
                run._r.insert(0, copy.deepcopy(label_run._r.rPr))
    if ph_values:
        for p in doc.paragraphs:
            _replace_in_paragraph(p, ph_values)
        for t in doc.tables:
            for row in t.rows:
                for c in _unique_cells(row):
                    for p in c.paragraphs:
                        _replace_in_paragraph(p, ph_values)
        for section in doc.sections:
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    _replace_in_paragraph(p, ph_values)
    doc.save(str(dst))


# ── Excel ─────────────────────────────────────────────────────────────


def _merged_top_left(ws) -> dict[tuple[int, int], tuple[int, int, int]]:
    """(行, 列) → (左上行, 左上列, 合并块最右列)。"""
    out = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = (rng.min_row, rng.min_col, rng.max_col)
    return out


def _xl_text(ws, r: int, c: int) -> str:
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def _inspect_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(str(path))
    slots: list[dict[str, Any]] = []
    has_placeholders = any(
        PLACEHOLDER_RE.search(str(c.value)) for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str)
    )
    for ws in wb.worksheets:
        merged = _merged_top_left(ws)
        maxr, maxc = ws.max_row, ws.max_column
        if has_placeholders:
            for r in range(1, maxr + 1):
                names = [(c, n) for c in range(1, maxc + 1)
                         for n in PLACEHOLDER_RE.findall(_xl_text(ws, r, c))]
                if len(names) >= 2:
                    slots.append({"kind": "table", "label": f"{ws.title} 第{r}行", "columns": [n for _, n in names],
                                  "loc": {"type": "xlsx_row_template", "sheet": ws.title, "row": r}})
                else:
                    slots += [{"kind": "text", "label": n, "loc": {"type": "placeholder", "name": n}} for _, n in names]
            continue
        used: set[tuple[int, int]] = set()
        # 1) 表头 + 空行
        for r in range(1, maxr + 1):
            cols = [c for c in range(1, maxc + 1) if _xl_text(ws, r, c) and merged.get((r, c), (r, c))[:2] == (r, c)]
            if len(cols) < 2 or not all(_is_label(_xl_text(ws, r, c)) for c in cols):
                continue
            # 表头那几格要挨着 (中间最多隔合并格), 不然是"标签 值 标签 值"的表单行
            if any(_xl_text(ws, r, c) == "" and (r, c) not in merged for c in range(cols[0], cols[-1] + 1)):
                continue
            empties = []
            span = range(cols[0], cols[-1] + 1)
            for rj in range(r + 1, maxr + 2):
                if rj > maxr or any(_xl_text(ws, rj, c) for c in span):
                    break
                # 第一行空行有边框 → 表格就是这几行带框的; 后面没框的空行是留白, 不算
                if empties and _xl_has_border(ws, empties[0], span) and not _xl_has_border(ws, rj, span):
                    break
                empties.append(rj)
            if empties:
                slots.append({"kind": "table", "label": _xl_title_above(ws, r, cols[0], used),
                              "columns": [_xl_text(ws, r, c) for c in cols],
                              "loc": {"type": "xlsx_rows", "sheet": ws.title, "header": r, "cols": cols, "rows": empties}})
                used |= {(rr, c) for rr in [r, *empties] for c in range(cols[0], cols[-1] + 1)}
        # 2) 标签格 + 右边 / 下边空格
        for r in range(1, maxr + 1):
            for c in range(1, maxc + 1):
                if (r, c) in used or merged.get((r, c), (r, c))[:2] != (r, c):
                    continue
                label = _xl_text(ws, r, c)
                if not _is_label(label):
                    continue
                right = merged.get((r, c), (r, c, c))[2] + 1
                below = merged.get((r + 1, c))
                alone = not any(_xl_text(ws, r, cc) for cc in range(1, maxc + 1) if cc != c)
                if (alone and below and below[:2] == (r + 1, c) and below[2] > c
                        and not _xl_text(ws, r + 1, c) and (r + 1, c) not in used):
                    # 「会议内容」独占一行, 下面是一整行合并的大空格 → 填下面 (右边那几格只是留白)
                    target = (r + 1, c)
                elif right <= maxc and not _xl_text(ws, r, right) and (r, right) not in used:
                    target = (r, right)
                elif r + 1 <= maxr and not _xl_text(ws, r + 1, c) and (r + 1, c) not in used:
                    target = (r + 1, c)
                else:
                    continue
                used |= {(r, c), target}
                slots.append({"kind": "text", "label": _clean_label(label),
                              "loc": {"type": "xlsx_cell", "sheet": ws.title, "row": target[0], "col": target[1]}})
    return _dedupe_placeholders(slots)


def _xl_title_above(ws, r: int, c: int, used: set | None = None) -> str:
    """多行表叫什么: 表头上面最近一格有字的 (「下一步工作」)。紧挨着的那格记进 used,
    不再把它当成"标签 + 右边空格"的单值空。"""
    for rr in range(r - 1, 0, -1):
        t = _xl_text(ws, rr, c)
        if t:
            if used is not None and rr == r - 1:
                used.add((rr, c))
            return _clean_label(t)[:_LABEL_MAX]
    return f"{ws.title} 表格"


def _xl_has_border(ws, r: int, cols: range) -> bool:
    return any(ws.cell(r, c).border.top.style or ws.cell(r, c).border.bottom.style for c in cols)


def _xl_copy_row_style(ws, src_row: int, dst_row: int, cols: range) -> None:
    for c in cols:
        s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
        if s.has_style:
            d._style = copy.copy(s._style)
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _xl_insert_rows(ws, at: int, n: int) -> None:
    """在 at 行前插 n 行; openpyxl 的 insert_rows 不挪合并单元格, 这里自己挪。"""
    moved = [rng for rng in ws.merged_cells.ranges if rng.min_row >= at]
    for rng in moved:
        ws.merged_cells.remove(rng)
    ws.insert_rows(at, n)
    for rng in moved:
        rng.shift(0, n)
        ws.merged_cells.add(rng)


def _xl_write(ws, r: int, c: int, text: str) -> None:
    cell = ws.cell(r, c)
    cell.value = text
    if "\n" in text:
        from openpyxl.styles import Alignment  # noqa: PLC0415
        a = copy.copy(cell.alignment)
        cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical or "top", wrap_text=True)


def _fill_xlsx(src: Path, dst: Path, slots: list[dict], values: dict[str, Any]) -> None:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(str(src))
    ph_values = {s["label"]: _as_text(values.get(s["id"], "")) for s in slots if s["loc"]["type"] == "placeholder"}
    # 从下往上写: 多行表插了行, 它下面的格子 (「记录人」之类) 行号会变, 先写下面的就不受影响
    def _row(s: dict) -> int:
        loc = s["loc"]
        return loc.get("row") or (loc.get("rows") or [0])[0]

    for s in sorted((s for s in slots if s["loc"]["type"] != "placeholder"), key=_row, reverse=True):
        loc, ws = s["loc"], wb[s["loc"]["sheet"]]
        if loc["type"] == "xlsx_cell":
            _xl_write(ws, loc["row"], loc["col"], _as_text(values.get(s["id"], "")))
            continue
        rows = _as_rows(values.get(s["id"]), len(s["columns"]))
        if loc["type"] == "xlsx_row_template":
            r0 = loc["row"]
            tmpl = {c: ws.cell(r0, c).value for c in range(1, ws.max_column + 1)}
            if len(rows) > 1:
                _xl_insert_rows(ws, r0 + 1, len(rows) - 1)
            for i, row in enumerate(rows or [[""] * len(s["columns"])]):
                r = r0 + i
                if i:
                    _xl_copy_row_style(ws, r0, r, range(1, ws.max_column + 1))
                vals = dict(zip(s["columns"], row))
                for c, v in tmpl.items():
                    if isinstance(v, str) and PLACEHOLDER_RE.search(v):
                        _xl_write(ws, r, c, PLACEHOLDER_RE.sub(lambda m, vals=vals: vals.get(m.group(1), ""), v))
        else:
            targets, cols = list(loc["rows"]), loc["cols"]
            extra = len(rows) - len(targets)
            if extra > 0:
                _xl_insert_rows(ws, targets[-1] + 1, extra)
                for k in range(extra):  # 照第一行空行 (表格本来的样子) 的格式
                    _xl_copy_row_style(ws, targets[0], targets[-1] + 1 + k, range(cols[0], cols[-1] + 1))
            for i, row in enumerate(rows):
                for c, text in zip(cols, row):
                    _xl_write(ws, targets[0] + i, c, text)
    if ph_values:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and PLACEHOLDER_RE.search(c.value):
                        c.value = PLACEHOLDER_RE.sub(lambda m: ph_values.get(m.group(1), m.group(0)), c.value)
    wb.save(str(dst))


# ── 共用 ─────────────────────────────────────────────────────────────


def _as_text(v: Any) -> str:
    if isinstance(v, list):
        return "\n".join(str(x) for x in v if str(x).strip())
    return "" if v is None else str(v).strip()


def _as_rows(v: Any, ncols: int) -> list[list[str]]:
    rows = []
    for row in v if isinstance(v, list) else []:
        cells = [str(x) if x is not None else "" for x in (row if isinstance(row, list) else [row])]
        rows.append((cells + [""] * ncols)[:ncols])
    return rows


def fill(src: Path, dst: Path, slots: list[dict], values: dict[str, Any]) -> None:
    if src.suffix.lower() == ".docx":
        _fill_docx(src, dst, slots, values)
    elif src.suffix.lower() == ".xlsx":
        _fill_xlsx(src, dst, slots, values)
    else:
        raise TemplateError(f"不支持的模版格式: {src.suffix}")


def slots_for_prompt(slots: list[dict]) -> list[dict]:
    """给大模型看的: 去掉文件里的定位信息。"""
    return [{k: s[k] for k in ("id", "label", "kind", "columns") if k in s} for s in slots]


# ── 待办表直接用结构化结果填 ──────────────────────────────────────────
#
# 10/3 真模型实测: 同一场会, Excel 模版的待办表填上了, Word 那次大模型给的是空 —— 表格
# 交给模型填不稳。待办本来就已经抽成结构化的了 (负责人 / 截止也是那一步定的), 认得出是
# 待办表就直接照抄, 不再经过模型。

_TASK_COL = re.compile(r"事项|任务|工作|内容|待办|措施")
_OWNER_COL = re.compile(r"负责|责任|承办|牵头|主办")
_DUE_COL = re.compile(r"时限|期限|截止|完成时间|时间|日期")
_SEQ_COL = re.compile(r"^(序号|编号|No\.?|#)$", re.IGNORECASE)


def is_todo_table(columns: list[str]) -> bool:
    has_task = any(_TASK_COL.search(c) for c in columns)
    return has_task and any(_OWNER_COL.search(c) or _DUE_COL.search(c) for c in columns)


def todo_rows(columns: list[str], action_items: list[dict]) -> list[list[str]]:
    rows = []
    for i, it in enumerate(action_items, 1):
        row = []
        for c in columns:
            if _SEQ_COL.match(c.strip()):
                row.append(str(i))
            elif _OWNER_COL.search(c):
                row.append(it.get("owner", ""))
            elif _DUE_COL.search(c):
                row.append(it.get("due", ""))
            elif _TASK_COL.search(c):
                row.append(it.get("task", ""))
            else:
                row.append("")
        rows.append(row)
    return rows


def apply_todos(slots: list[dict], values: dict, action_items: list[dict]) -> dict:
    """认得出是待办表的, 用结构化待办覆盖模型给的值 (没抽到待办就保留模型的)。"""
    if not action_items:
        return values
    out = dict(values)
    for s in slots:
        if s["kind"] == "table" and is_todo_table(s.get("columns") or []):
            out[s["id"]] = todo_rows(s["columns"], action_items)
    return out
