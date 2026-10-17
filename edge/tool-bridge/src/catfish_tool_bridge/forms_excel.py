"""一起填表 · Excel: 从模板建表 / 汇总导出 (docs/FORM-COLLECT-PLAN.md §2 §4)。

建表: 第一行做表头; 每列推断类型 —— 数据验证的下拉列表 → 单选, 第二行单元格是日期格式 → 日期,
数字格式 → 数字; 推不出看列名 (「日期」→ 日期, 「次数 / 数量 / 金额」→ 数字, 「说明 / 备注」→ 长文本),
都不像就是文本。这只是起点, 员工在界面上改得了。

导出: 有模板就打开模板, 从表头下第一个空行写, 每个字段写回它原来那一列; 最后加「填写人 / 提交时间 / 状态」。
没模板用内置样式。另有一页「进度」列每个收件人的状态。
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

from . import advisor_io, forms

_DATE_WORDS = ("日期", "时间", "截止", "生日")
_NUMBER_WORDS = ("次数", "数量", "金额", "人数", "个数", "件数", "天数", "台数", "得分", "分数", "合计", "总数", "费用")
_LONG_WORDS = ("说明", "备注", "情况", "描述", "原因", "意见", "建议", "措施", "总结")


def _list_options(ws, col_letter: str) -> list[str]:
    """该列 (第 2 行起) 套了「序列」数据验证就返回选项。只认直接写在公式里的 "a,b,c"。"""
    for dv in getattr(ws.data_validations, "dataValidation", []) or []:
        if dv.type != "list" or not dv.formula1:
            continue
        if not any(col_letter == re.sub(r"\d", "", str(rng).split(":")[0]) for rng in str(dv.sqref).split()):
            continue
        f = str(dv.formula1).strip()
        if f.startswith('"') and f.endswith('"'):
            opts = [o.strip() for o in f[1:-1].split(",") if o.strip()]
            if len(opts) >= 2:
                return opts
    return []


def _guess_type(label: str, sample) -> str:
    fmt = (getattr(sample, "number_format", "") or "").lower()
    if getattr(sample, "is_date", False) or any(t in fmt for t in ("yy", "mm-dd", "m/d", "d/m")):
        return "date"
    if fmt not in ("", "general", "@") and re.search(r"[0#]", fmt):
        return "number"
    if any(w in label for w in _DATE_WORDS):
        return "date"
    if any(w in label for w in _NUMBER_WORDS):
        return "number"
    if any(w in label for w in _LONG_WORDS):
        return "longtext"
    return "text"


def infer_fields(path: str) -> dict[str, Any]:
    """读 Excel 模板第一行 → 字段建议 (不落盘)。"""
    from openpyxl import load_workbook  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    p = Path(path).expanduser()
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("只支持 .xlsx / .xlsm (老的 .xls 先用 Excel 另存一下)")
    if not p.is_file():
        raise ValueError(f"找不到文件: {p}")
    wb = load_workbook(p)
    ws = wb.worksheets[0]
    fields = []
    for cell in ws[1]:
        label = str(cell.value or "").strip()
        if not label:
            continue
        letter = get_column_letter(cell.column)
        opts = _list_options(ws, letter)
        f: dict[str, Any] = {"label": label, "col": cell.column,
                             "type": "choice" if opts else _guess_type(label, ws.cell(row=2, column=cell.column))}
        if opts:
            f["options"] = opts
        fields.append(f)
    if not fields:
        raise ValueError("第一行没有表头 (表头要放在第一个工作表的第一行)")
    return {"ok": True, "title": p.stem, "fields": forms.normalize_fields(fields), "sheet": ws.title}


def _cell_value(field: dict[str, Any], v):
    if v is None:
        return None
    if field["type"] == "date":
        try:
            return date.fromisoformat(str(v))
        except ValueError:
            return str(v)
    if field["type"] == "multichoice" and isinstance(v, list):
        return "、".join(v)
    return v


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_")[:60] or "表单"


def export(form_id: str) -> dict[str, Any]:
    from openpyxl import Workbook, load_workbook  # noqa: PLC0415
    from openpyxl.styles import Font  # noqa: PLC0415

    form = forms.load_outgoing(form_id)
    resp = forms.responses_of(form_id)
    prog = forms.progress(form)
    state_of = {r["email"]: r["state"] for r in prog["rows"]}
    tpl = forms._out_dir(form_id) / (form.get("template") or "") if form.get("template") else None
    fields = form["fields"]

    if tpl and tpl.is_file():
        wb = load_workbook(tpl)
        ws = wb.worksheets[0]
        cols = {f["key"]: f.get("col") for f in fields}
        last = max([c for c in cols.values() if c] + [ws.max_column])
        nxt = last + 1
        for f in fields:            # 模板里没有原列的字段 (界面上后加的) 接在后面
            if not cols.get(f["key"]):
                cols[f["key"]] = nxt
                ws.cell(row=1, column=nxt, value=f["label"])
                nxt += 1
        row = 2
        while any(ws.cell(row=row, column=c).value not in (None, "") for c in cols.values() if c):
            row += 1
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = _safe(form["title"])[:31]
        cols = {f["key"]: i for i, f in enumerate(fields, 1)}
        for f in fields:
            ws.cell(row=1, column=cols[f["key"]], value=f["label"]).font = Font(bold=True)
        nxt, row = len(fields) + 1, 2
    extra = {"填写人": nxt, "提交时间": nxt + 1, "状态": nxt + 2}
    for label, c in extra.items():
        ws.cell(row=1, column=c, value=label).font = Font(bold=True)

    written = 0
    for email, r in resp.items():
        if r.get("status") != "submitted":
            continue
        vals = r.get("values") or {}
        for f in fields:
            ws.cell(row=row, column=cols[f["key"]], value=_cell_value(f, vals.get(f["key"])))
        ws.cell(row=row, column=extra["填写人"], value=email)
        ws.cell(row=row, column=extra["提交时间"], value=(r.get("at") or "")[:16].replace("T", " "))
        ws.cell(row=row, column=extra["状态"], value=state_of.get(email, "已提交"))
        row += 1
        written += 1

    ps = wb.create_sheet("进度")
    for c, h in enumerate(("收件人", "状态", "提交时间", "不填原因", "催过"), 1):
        ps.cell(row=1, column=c, value=h).font = Font(bold=True)
    for i, r in enumerate(prog["rows"], 2):
        ps.cell(row=i, column=1, value=r["email"])
        ps.cell(row=i, column=2, value=r["state"])
        ps.cell(row=i, column=3, value=(r.get("at") or "")[:16].replace("T", " "))
        ps.cell(row=i, column=4, value=r.get("reason") or "")
        ps.cell(row=i, column=5, value="是" if r.get("reminded_at") else "")

    path = advisor_io.today_outputs_dir() / f"{_safe(form['title'])}-汇总.xlsx"
    wb.save(path)
    return {"ok": True, "path": str(path), "rows": written, "progress": prog["counts"]}
