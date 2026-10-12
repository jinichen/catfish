"""出报价单 (销售跟进阶段 2, docs/SALES-FOLLOWUP-PLAN.md §10.2)。

catfish_customer action=quote: 客户 + 明细 → ~/.catfish/outputs/<日期>/报价-<客户>-v<N>.xlsx。

- 单价默认取价格表 (sales_pricebook) 标价; 价格表里没有的商品必须给 price, 不猜价。
- 折后单价低于价格表「最低价」只出提议 (needs_confirmation), 员工确认才生成。
- 版本号按这个客户已出过的报价自动 +1, 记在 ~/.catfish/sales/quotes.json, 周报从这里数。
- 生成后往客户时间线记一条「系统」行 (未发出); 真发给客户后由小鲶 log kind=报价, 雷达的
  「报价没回音」从那条开始算 —— 生成 ≠ 发出, 不替销售发。
- 有公司模板 (quote_template.xlsx) 就按模板填, 否则用内置版式。
"""
from __future__ import annotations

import json
import re
from copy import copy
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from . import advisor_io, customers, sales_pricebook

_DIGITS = "零壹贰叁肆伍陆柒捌玖"
_UNITS = ["", "拾", "佰", "仟"]
_SECTIONS = ["", "万", "亿", "兆"]


def rmb_upper(amount: float) -> str:
    """人民币大写: 12345.6 → 壹万贰仟叁佰肆拾伍元陆角整。"""
    cents = int(round(abs(amount) * 100))
    yuan, jiao, fen = cents // 100, cents // 10 % 10, cents % 10
    if yuan == 0:
        int_part = ""
    else:
        int_part, sec_i, zero_pending = "", 0, False
        while yuan:
            sec, yuan = yuan % 10000, yuan // 10000
            s, zero = "", False
            for i in range(4):
                d = sec // (10 ** i) % 10
                if d == 0:
                    zero = bool(s)
                else:
                    s = _DIGITS[d] + _UNITS[i] + ("零" if zero else "") + s
                    zero = False
            if sec:
                int_part = s + _SECTIONS[sec_i] + ("零" if zero_pending else "") + int_part
                zero_pending = sec < 1000
            else:
                zero_pending = bool(int_part)
            sec_i += 1
        int_part += "元"
    tail = ""
    if jiao:
        tail += _DIGITS[jiao] + "角"
    if fen:
        tail += ("零" if not jiao and int_part else "") + _DIGITS[fen] + "分"
    out = (int_part + tail) or "零元"
    return ("负" if amount < 0 else "") + out + ("" if fen else "整")


def _ledger_path() -> Path:
    return sales_pricebook.sales_dir() / "quotes.json"


def load_ledger() -> list[dict[str, Any]]:
    try:
        rows = json.loads(_ledger_path().read_text(encoding="utf-8"))
        return rows if isinstance(rows, list) else []
    except (OSError, ValueError):
        return []


def _save_ledger(rows: list[dict[str, Any]]) -> None:
    tmp = _ledger_path().with_suffix(".json.tmp")
    tmp.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(_ledger_path())


def _discount(v: Any) -> float:
    """0.9 = 九折; 也认 9 / 8.5 这种「几折」写法。不给 = 不打折。"""
    if v in (None, ""):
        return 1.0
    d = sales_pricebook.parse_money(str(v).replace("折", ""))
    if d is None or d <= 0:
        raise ValueError(f"折扣写法认不出: {v} (0.9 = 九折)")
    if d > 1:
        d /= 10
    if d > 1:
        raise ValueError(f"折扣不能大于 1: {v}")
    return d


def price_lines(raw_items: Any) -> list[dict[str, Any]]:
    """把模型给的明细对上价格表, 算出每行金额。对不上价的直接报错 + 候选。"""
    if not isinstance(raw_items, list) or not raw_items:
        raise ValueError("items 至少一行: [{name, spec?, qty, price?, discount?, note?}]")
    lines = []
    for i, it in enumerate(raw_items, 1):
        if not isinstance(it, dict) or not str(it.get("name") or "").strip():
            raise ValueError(f"第 {i} 行缺商品名")
        name, spec = str(it["name"]).strip(), str(it.get("spec") or "").strip()
        qty = sales_pricebook.parse_money(it.get("qty", 1))
        if qty is None or qty <= 0:
            raise ValueError(f"「{name}」数量不对: {it.get('qty')}")
        row, candidates = sales_pricebook.lookup(name, spec)
        given = sales_pricebook.parse_money(it.get("price")) if it.get("price") not in (None, "") else None
        if row is None and given is None:
            hint = f"价格表里相近的: {candidates}" if candidates else "价格表里没有这个商品"
            raise ValueError(f"「{name}{' / ' + spec if spec else ''}」{hint}。用价格表里的准确名字, 或者问员工单价后带 price")
        list_price = row["price"] if row else given
        unit_price = given if given is not None else list_price
        disc = _discount(it.get("discount"))
        final_unit = round(unit_price * disc, 2)
        lines.append({
            "name": row["name"] if row else name,
            "spec": (row["spec"] if row else spec) or spec,
            "unit": str(it.get("unit") or (row["unit"] if row else "")).strip(),
            "qty": qty,
            "list_price": list_price,
            "unit_price": unit_price,
            "discount": disc,
            "final_unit": final_unit,
            "amount": round(final_unit * qty, 2),
            "floor": row.get("floor") if row else None,
            "in_pricebook": row is not None,
            "note": str(it.get("note") or "").strip(),
        })
    return lines


def _fmt_qty(q: float) -> int | float:
    return int(q) if float(q).is_integer() else q


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_") or "客户"


# ── 写 xlsx ──────────────────────────────────────────────────────────────

def _write_builtin(path: Path, q: dict[str, Any]) -> None:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    widths = [8, 22, 12, 7, 7, 12, 6, 13, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.merge_cells("A1:I1")
    ws["A1"] = f"{q['seller']['company']} 报价单" if q["seller"]["company"] else "报价单"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    meta = [("客户", q["customer"], "报价编号", q["quote_no"]),
            ("联系人", q["customer_contact"], "报价日期", q["date"]),
            ("报价方", q["seller"]["company"], "有效期至", q["valid_until"])]
    for r, (k1, v1, k2, v2) in enumerate(meta, 2):
        ws.cell(r, 1, k1).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(r, 2, v1)
        ws.cell(r, 6, k2).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
        ws.cell(r, 7, v2)

    hr = 6
    for c, h in enumerate(["序号", "商品", "规格", "单位", "数量", "单价", "折扣", "金额", "备注"], 1):
        cell = ws.cell(hr, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True), head_fill, box
        cell.alignment = Alignment(horizontal="center")
    r = hr
    for i, ln in enumerate(q["lines"], 1):
        r += 1
        vals = [i, ln["name"], ln["spec"], ln["unit"], _fmt_qty(ln["qty"]), ln["unit_price"],
                "" if ln["discount"] == 1 else f"{ln['discount'] * 10:g}折", ln["amount"], ln["note"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = box
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c in (6, 8):
                cell.number_format = "#,##0.00"
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(r, 1, f"合计 (大写): {q['total_upper']}").font = Font(bold=True)
    total = ws.cell(r, 8, q["total"])
    total.font, total.number_format = Font(bold=True), "#,##0.00"
    for c in range(1, 10):
        ws.cell(r, c).border = box

    r += 2
    if q["note"]:
        ws.cell(r, 1, "备注").font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.cell(r, 2, q["note"]).alignment = Alignment(wrap_text=True)
        r += 1
    if q["seller"]["terms"]:
        ws.cell(r, 1, "条款").font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.cell(r, 2, q["seller"]["terms"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (q["seller"]["terms"].count("\n") + 1)
        r += 1
    contact = " · ".join(x for x in (q["seller"]["contact"], q["seller"]["phone"], q["seller"]["email"]) if x)
    if contact:
        r += 1
        ws.cell(r, 1, "联系").font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.cell(r, 2, contact)
    _fit_width(ws)
    wb.save(path)


def _fit_width(ws: Any) -> None:
    """打印 / 转 PDF 时缩到一页宽, 金额列不会被挤到第二页。"""
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


_PH = re.compile(r"\{\{(.+?)\}\}")


def _fill(cell: Any, values: dict[str, Any]) -> None:
    if not isinstance(cell.value, str) or "{{" not in cell.value:
        return
    whole = _PH.fullmatch(cell.value.strip())
    if whole and whole.group(1) in values:
        cell.value = values[whole.group(1)]  # 整格一个占位: 保留数字类型, Excel 里还能算
        return
    cell.value = _PH.sub(lambda m: str(values.get(m.group(1), m.group(0))), cell.value).strip()  # 「{{商品}} {{规格}}」规格空时不留尾巴


def _write_custom(path: Path, tpl: Path, q: dict[str, Any]) -> None:
    """按公司模板填: 有 {{商品}} 的那一行是明细行, 按条数往下复制; 其它格填抬头 / 合计。"""
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(tpl)
    scalars = {"客户": q["customer"], "客户联系人": q["customer_contact"], "日期": q["date"], "报价日期": q["date"],
               "报价编号": q["quote_no"], "版本": f"v{q['version']}", "有效期至": q["valid_until"],
               "合计": q["total"], "合计大写": q["total_upper"], "备注": q["note"],
               "报价方": q["seller"]["company"], "联系人": q["seller"]["contact"], "电话": q["seller"]["phone"],
               "邮箱": q["seller"]["email"], "条款": q["seller"]["terms"]}
    for ws in wb.worksheets:
        item_row = next((r for r in range(1, ws.max_row + 1)
                         if any(isinstance(c.value, str) and "{{商品}}" in c.value for c in ws[r])), None)
        if item_row is not None:
            n = len(q["lines"])
            if n > 1:
                # openpyxl 插行不挪合并单元格, 自己挪
                below = [m for m in ws.merged_cells.ranges if m.min_row > item_row]
                for m in below:
                    ws.unmerge_cells(str(m))
                ws.insert_rows(item_row + 1, n - 1)
                for m in below:
                    m.shift(0, n - 1)
                    ws.merge_cells(str(m))
                height = ws.row_dimensions[item_row].height
                for k in range(1, n):
                    for c in ws[item_row]:
                        dst = ws.cell(item_row + k, c.column, c.value)
                        if c.has_style:
                            dst._style = copy(c._style)
                    if height:
                        ws.row_dimensions[item_row + k].height = height
            for k, ln in enumerate(q["lines"]):
                vals = {"序号": k + 1, "商品": ln["name"], "规格": ln["spec"], "单位": ln["unit"],
                        "数量": _fmt_qty(ln["qty"]), "单价": ln["unit_price"], "标价": ln["list_price"],
                        "折扣": "" if ln["discount"] == 1 else f"{ln['discount'] * 10:g}折",
                        "折后单价": ln["final_unit"], "金额": ln["amount"], "备注": ln["note"]}
                for c in ws[item_row + k]:
                    _fill(c, vals)
        for row in ws.iter_rows():
            for c in row:
                _fill(c, scalars)
    wb.save(path)


# ── action=quote ─────────────────────────────────────────────────────────

def action_quote(args: dict[str, Any]) -> dict[str, Any]:
    rec = customers._resolve(args)
    lines = price_lines(args.get("items"))
    below = [ln for ln in lines if ln["floor"] is not None and ln["final_unit"] < ln["floor"]]
    total = round(sum(ln["amount"] for ln in lines), 2)
    if below and not args.get("confirmed"):
        return {"ok": False, "needs_confirmation": True,
                "below_floor": [{"name": ln["name"], "spec": ln["spec"], "final_unit": ln["final_unit"], "floor": ln["floor"]}
                                for ln in below],
                "total": total,
                "message": "有商品折后单价低于价格表最低价: 跟员工说清楚 (哪个商品、折后多少、最低价多少), "
                           "员工确认 (必要时先找主管) 后带 confirmed=true 再调"}

    settings = sales_pricebook.load_quote_settings()
    today = date.today()
    ledger = load_ledger()
    version = 1 + max((x.get("version", 0) for x in ledger if x.get("customer") == rec["name"]), default=0)
    seq = 1 + sum(1 for x in ledger if x.get("date") == today.isoformat())
    valid_days = int(args.get("valid_days") or settings["valid_days"])
    contact = next((x.get("name") for x in rec["contacts"] if x.get("name")), "")
    q = {
        "customer": rec["name"], "customer_contact": contact, "version": version,
        "quote_no": f"Q{today:%Y%m%d}-{seq:02d}", "date": today.isoformat(),
        "valid_until": (today + timedelta(days=valid_days)).isoformat(),
        "lines": lines, "total": total, "total_upper": rmb_upper(total),
        "note": str(args.get("note") or "").strip(),
        "seller": {**settings, "contact": settings["contact"] or str(rec.get("owner") or "")},
    }
    path = advisor_io.today_outputs_dir() / f"报价-{_safe(rec['name'])}-v{version}.xlsx"
    tpl = sales_pricebook.quote_template_path()
    if tpl.is_file():
        _write_custom(path, tpl, q)
    else:
        _write_builtin(path, q)

    ledger.append({"customer": rec["name"], "version": version, "quote_no": q["quote_no"], "date": q["date"],
                   "total": total, "file": str(path), "valid_until": q["valid_until"],
                   "items": [{k: ln[k] for k in ("name", "spec", "qty", "final_unit", "amount")} for ln in lines]})
    _save_ledger(ledger)
    customers.append_system(rec, f"生成报价单 v{version} 合计 ¥{total:,.2f} (未发出) {path.name}")
    return {
        "ok": True, "file": str(path), "customer": rec["name"], "version": version, "quote_no": q["quote_no"],
        "total": total, "total_upper": q["total_upper"], "valid_until": q["valid_until"],
        "lines": [{k: ln[k] for k in ("name", "spec", "qty", "unit_price", "discount", "amount", "in_pricebook")} for ln in lines],
        "below_floor_confirmed": bool(below),
        "next": "把文件路径给员工核对; 员工发给客户之后, 调 action=log who=我方 kind=报价 "
                f"text=「v{version} 合计 ¥{total:,.2f}」记一笔 (雷达从那天开始算报价回音)。不要替员工发。",
    }


def action_quotes(args: dict[str, Any]) -> dict[str, Any]:
    """这个客户出过的报价版本 (新的在前)。"""
    rec = customers._resolve(args)
    rows = [x for x in load_ledger() if x.get("customer") == rec["name"]]
    return {"ok": True, "customer": rec["name"], "quotes": sorted(rows, key=lambda x: x.get("version", 0), reverse=True)}
