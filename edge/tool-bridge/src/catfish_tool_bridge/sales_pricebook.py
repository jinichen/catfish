"""价格表 + 报价单设置 (销售跟进阶段 2, docs/SALES-FOLLOWUP-PLAN.md §10)。

价格表 = 销售团队现有的 Excel, 在设置「🤝 销售跟进」里导入。表头不要求固定写法, 按常见叫法认列
(商品 / 规格 / 单位 / 标价 / 最低价 / 备注), 认不出「商品」和「标价」两列就报错, 把表头念给员工。
导入后存成 ~/.catfish/sales/pricebook.json, 原 Excel 留一份 pricebook_source.xlsx 方便对照。

报价单设置 (我方公司名、联系人、默认有效期、条款) 和公司自己的报价单模板也放在 ~/.catfish/sales/。
不随商品写死, 不同团队各导各的。
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from . import wiki_files

# 列名 → 常见叫法 (表头去空格后完全相等才算; 先列的优先)
COLUMNS: dict[str, tuple[str, ...]] = {
    "name": ("商品", "商品名称", "产品", "产品名称", "名称", "品名", "项目", "服务", "服务项目"),
    "spec": ("规格", "型号", "规格型号", "版本", "配置"),
    "unit": ("单位", "计价单位"),
    "price": ("标价", "单价", "价格", "报价", "售价", "指导价", "零售价", "目录价", "单价(元)", "价格(元)", "标价(元)"),
    "floor": ("最低价", "底价", "最低售价", "限价", "最低价(元)", "底价(元)"),
    "note": ("备注", "说明"),
}
REQUIRED = ("name", "price")

DEFAULT_QUOTE_SETTINGS: dict[str, Any] = {
    "company": "",   # 我方公司名 (报价方)
    "contact": "",   # 我方联系人, 空 = 用客户页的负责销售
    "phone": "",
    "email": "",
    "valid_days": 30,
    "terms": "1. 以上报价为人民币含税价。\n2. 报价有效期内价格不变, 逾期请重新确认。",
}


def sales_dir() -> Path:
    d = wiki_files._catfish_home() / "sales"
    d.mkdir(parents=True, exist_ok=True)
    return d


def pricebook_path() -> Path:
    return sales_dir() / "pricebook.json"


def quote_settings_path() -> Path:
    return sales_dir() / "quote_settings.json"


def quote_template_path() -> Path:
    return sales_dir() / "quote_template.xlsx"


def _atomic_write(path: Path, data: Any) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def parse_money(v: Any) -> float | None:
    """「¥1,200」「1.2万」「12000.00」→ 数; 认不出返回 None。"""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[¥￥,，\s元]", "", str(v))
    m = re.fullmatch(r"(\d+(?:\.\d+)?)(万)?", s)
    if not m:
        return None
    return float(m.group(1)) * (10000 if m.group(2) else 1)


def _norm(h: Any) -> str:
    """表头归一: 去空白, 全角括号转半角 (「单价（元）」=「单价(元)」)。"""
    return re.sub(r"\s", "", str(h or "")).replace("（", "(").replace("）", ")")


def _map_header(row: tuple[Any, ...]) -> dict[str, int]:
    cells = [_norm(c) for c in row]
    mapping: dict[str, int] = {}
    for key, names in COLUMNS.items():
        for n in names:
            if n in cells and cells.index(n) not in mapping.values():
                mapping[key] = cells.index(n)
                break
    return mapping


def _open(load_workbook: Any, src: Path, **kw: Any) -> Any:
    try:
        return load_workbook(src, **kw)
    except Exception as e:  # noqa: BLE001 —— 坏 zip / 加密 / 其实是 xls 改了后缀, 都归成「打不开」
        raise ValueError(f"打不开这个 Excel ({type(e).__name__}); 在 Excel 里另存为 .xlsx 再试") from None


# ── 价格表 ───────────────────────────────────────────────────────────────

def import_pricebook(path: str) -> dict[str, Any]:
    """读 Excel 价格表 (第一个认得出表头的 sheet, 表头在前 10 行内)。"""
    from openpyxl import load_workbook  # noqa: PLC0415

    src = Path(path).expanduser()
    if src.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("价格表要 .xlsx (老的 .xls 请在 Excel 里另存为 .xlsx)")
    if not src.is_file():
        raise ValueError(f"找不到文件: {src}")
    wb = _open(load_workbook, src, read_only=True, data_only=True)
    seen_headers: list[str] = []
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for hi, row in enumerate(rows[:10]):
                mapping = _map_header(row)
                if all(k in mapping for k in REQUIRED):
                    break
                if any(c not in (None, "") for c in row) and not seen_headers:
                    seen_headers = [str(c) for c in row if c not in (None, "")]
            else:
                continue
            items, skipped = [], 0
            for row in rows[hi + 1:]:
                get = lambda k: row[mapping[k]] if k in mapping and mapping[k] < len(row) else None  # noqa: E731
                name = str(get("name") or "").strip()
                price = parse_money(get("price"))
                if not name:
                    continue
                if price is None:
                    skipped += 1
                    continue
                items.append({
                    "name": name,
                    "spec": str(get("spec") or "").strip(),
                    "unit": str(get("unit") or "").strip(),
                    "price": price,
                    "floor": parse_money(get("floor")),
                    "note": str(get("note") or "").strip(),
                })
            if not items:
                raise ValueError(f"「{ws.title}」认出了表头, 但下面没有一行带价格的商品")
            book = {"source": src.name, "sheet": ws.title, "imported_at": datetime.now().isoformat(timespec="seconds"),
                    "columns": {k: _norm(rows[hi][i]) for k, i in mapping.items()}, "items": items}
            _atomic_write(pricebook_path(), book)
            shutil.copyfile(src, sales_dir() / "pricebook_source.xlsx")
            return {**book, "count": len(items), "skipped": skipped}
    finally:
        wb.close()
    raise ValueError(f"没认出「商品」和「价格」两列。第一行表头是: {seen_headers or '空'}。"
                     f"把商品列叫「商品 / 产品 / 名称」, 价格列叫「标价 / 单价 / 价格」之一再导入。")


def load_pricebook() -> dict[str, Any]:
    try:
        book = json.loads(pricebook_path().read_text(encoding="utf-8"))
        if isinstance(book, dict) and isinstance(book.get("items"), list):
            return book
    except (OSError, ValueError):
        pass
    return {"source": None, "items": []}


def search(query: str, limit: int = 20) -> list[dict[str, Any]]:
    q = query.strip().lower()
    items = load_pricebook()["items"]
    if not q:
        return items[:limit]
    return [x for x in items if q in x["name"].lower() or q in x["spec"].lower()][:limit]


def lookup(name: str, spec: str = "") -> tuple[dict[str, Any] | None, list[str]]:
    """报价时找价格表里的那一行: 名字 (+规格) 完全对上才算; 对不上返回候选, 不猜。"""
    name, spec = name.strip(), spec.strip()
    items = load_pricebook()["items"]
    hits = [x for x in items if x["name"] == name and (not spec or x["spec"] == spec)]
    if len(hits) == 1:
        return hits[0], []
    if not hits:
        hits = [x for x in items if name and (name in x["name"] or x["name"] in name)]
    return None, [f"{x['name']}{' / ' + x['spec'] if x['spec'] else ''}" for x in hits[:10]]


# ── 报价单设置 / 模板 ─────────────────────────────────────────────────────

def load_quote_settings() -> dict[str, Any]:
    out = dict(DEFAULT_QUOTE_SETTINGS)
    try:
        user = json.loads(quote_settings_path().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        user = {}
    if isinstance(user, dict):
        for k, v in user.items():
            if k == "valid_days" and isinstance(v, (int, float)) and v > 0:
                out[k] = int(v)
            elif k in DEFAULT_QUOTE_SETTINGS and k != "valid_days" and isinstance(v, str):
                out[k] = v.strip()
    out["custom_template"] = quote_template_path().is_file()
    return out


def save_quote_settings(s: dict[str, Any]) -> dict[str, Any]:
    clean = {k: s[k] for k in DEFAULT_QUOTE_SETTINGS if k in s}
    if "valid_days" in clean:
        try:
            clean["valid_days"] = max(1, int(clean["valid_days"]))
        except (TypeError, ValueError):
            raise ValueError("有效期天数要是正整数") from None
    _atomic_write(quote_settings_path(), clean)
    return load_quote_settings()


def import_quote_template(path: str) -> dict[str, Any]:
    """公司自己的报价单 .xlsx: 单元格里放 {{客户}} {{合计}} 这类占位, 明细行放 {{商品}} {{数量}} …"""
    from openpyxl import load_workbook  # noqa: PLC0415

    src = Path(path).expanduser()
    if src.suffix.lower() != ".xlsx" or not src.is_file():
        raise ValueError("报价单模板要 .xlsx 文件")
    wb = _open(load_workbook, src)
    found = sorted({m for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) for m in re.findall(r"\{\{(.+?)\}\}", c.value)})
    if "商品" not in found:
        raise ValueError("模板里没有 {{商品}} 占位 —— 明细那一行要有 {{商品}} {{数量}} {{单价}} {{金额}} 这类占位, 小鲶才知道往哪填")
    shutil.copyfile(src, quote_template_path())
    return {"ok": True, "placeholders": found, **load_quote_settings()}


def clear_quote_template() -> dict[str, Any]:
    quote_template_path().unlink(missing_ok=True)
    return load_quote_settings()
