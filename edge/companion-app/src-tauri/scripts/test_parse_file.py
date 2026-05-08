"""parse_file.py preview-only mode 单测 (5/5 重构).

跑法: pytest src-tauri/scripts/test_parse_file.py
依赖: openpyxl, pypdfium2, python-docx (上传支持的格式)
"""
from __future__ import annotations

import csv
import json
import subprocess
import sys
from pathlib import Path

import pytest

THIS = Path(__file__).resolve()
PARSE_PY = THIS.parent / "parse_file.py"


def run_parse(file_path: Path) -> dict:
    """跑 parse_file.py 返 JSON dict (或 raise)."""
    res = subprocess.run(
        [sys.executable, str(PARSE_PY), str(file_path)],
        capture_output=True, text=True,
    )
    if res.returncode != 0:
        raise RuntimeError(f"parse_file.py 退出 {res.returncode}: stderr={res.stderr}")
    return json.loads(res.stdout)


# ============================================================
# Text (txt/md/log) — 最简单, 不需要外部依赖
# ============================================================

def test_text_short(tmp_path: Path) -> None:
    """短文本, 不超 PREVIEW_MAX_CHARS, 全文返回."""
    f = tmp_path / "short.txt"
    f.write_text("hello world", encoding="utf-8")
    r = run_parse(f)
    assert r["kind"] == "text"
    assert r["filename"] == "short.txt"
    assert r["meta"]["total_chars"] == 11
    assert "hello world" in r["preview_text"]


def test_text_long(tmp_path: Path) -> None:
    """长文本, 超 PREVIEW_MAX_CHARS (5000), 截 preview + 提示用 execute_code."""
    f = tmp_path / "long.txt"
    content = "a" * 8000  # 8K, 超 5K limit
    f.write_text(content, encoding="utf-8")
    r = run_parse(f)
    assert r["kind"] == "text"
    assert r["meta"]["total_chars"] == 8000
    assert "execute_code" in r["preview_text"]
    assert r["preview_chars"] < 8000  # preview 比原文小


def test_md_kind_is_text(tmp_path: Path) -> None:
    """.md 也归类 text."""
    f = tmp_path / "doc.md"
    f.write_text("# Hello\n\nWorld", encoding="utf-8")
    r = run_parse(f)
    assert r["kind"] == "text"
    assert r["ext"] == ".md"


# ============================================================
# CSV
# ============================================================

def test_csv_small(tmp_path: Path) -> None:
    """小 CSV, 全部进 preview."""
    f = tmp_path / "data.csv"
    with f.open("w", encoding="utf-8", newline="") as out:
        w = csv.writer(out)
        w.writerow(["name", "age"])
        w.writerow(["alice", "30"])
        w.writerow(["bob", "25"])
    r = run_parse(f)
    assert r["kind"] == "csv"
    assert r["meta"]["total_rows"] == 3
    assert "alice" in r["preview_text"]
    assert "bob" in r["preview_text"]


def test_csv_truncated_large(tmp_path: Path) -> None:
    """100 行 CSV, preview 只显前 N 行 + 总行数提示."""
    f = tmp_path / "big.csv"
    with f.open("w", encoding="utf-8", newline="") as out:
        w = csv.writer(out)
        w.writerow(["id", "name"])
        for i in range(100):
            w.writerow([str(i), f"name_{i}"])
    r = run_parse(f)
    assert r["kind"] == "csv"
    assert r["meta"]["total_rows"] == 101  # 100 + header
    assert "execute_code" in r["preview_text"]
    # 前 N 行肯定在 preview 里
    assert "name_0" in r["preview_text"]
    # 第 90 行多半不在 (PREVIEW_ROWS=20)
    assert "name_90" not in r["preview_text"]


# ============================================================
# Excel (5/5 鸿波核心场景)
# ============================================================

@pytest.fixture
def has_openpyxl() -> bool:
    try:
        import openpyxl  # noqa: F401, PLC0415
        return True
    except ImportError:
        pytest.skip("openpyxl 未装")


def test_excel_multi_sheet(tmp_path: Path, has_openpyxl: bool) -> None:
    """多 sheet Excel, preview 列每 sheet 列头 + 前 N 行 + 总行数."""
    from openpyxl import Workbook  # noqa: PLC0415
    f = tmp_path / "multi.xlsx"
    wb = Workbook()
    s1 = wb.active
    s1.title = "1月"
    s1.append(["date", "revenue"])
    for d in range(1, 32):
        s1.append([f"2026-01-{d:02d}", 1000 + d])
    s2 = wb.create_sheet("2月")
    s2.append(["date", "revenue"])
    for d in range(1, 29):
        s2.append([f"2026-02-{d:02d}", 1500 + d])
    wb.save(f)

    r = run_parse(f)
    assert r["kind"] == "excel"
    assert r["meta"]["total_sheets"] == 2
    assert r["meta"]["sheets"] == ["1月", "2月"]
    # 列头在 preview
    assert "date" in r["preview_text"]
    assert "revenue" in r["preview_text"]
    # sheet 名标记在 preview
    assert "## Sheet: 1月" in r["preview_text"]
    assert "## Sheet: 2月" in r["preview_text"]
    # 总行数提示在
    assert "execute_code" in r["preview_text"]


def test_excel_row_counts(tmp_path: Path, has_openpyxl: bool) -> None:
    """row_counts meta 字段正确."""
    from openpyxl import Workbook  # noqa: PLC0415
    f = tmp_path / "counts.xlsx"
    wb = Workbook()
    s = wb.active
    s.title = "Sheet1"
    s.append(["x"])
    for i in range(50):
        s.append([i])
    wb.save(f)

    r = run_parse(f)
    counts = r["meta"]["row_counts"]
    # 51 = 1 header + 50 data
    assert counts["Sheet1"] == 51


# ============================================================
# 错误路径
# ============================================================

def test_unsupported_ext(tmp_path: Path) -> None:
    """不支持的扩展名 → error JSON, 退出码 4."""
    f = tmp_path / "x.bin"
    f.write_bytes(b"\x00\x01\x02")
    res = subprocess.run(
        [sys.executable, str(PARSE_PY), str(f)],
        capture_output=True, text=True,
    )
    assert res.returncode == 4
    err = json.loads(res.stdout)
    assert "不支持" in err["error"]


def test_missing_file() -> None:
    """文件不存在 → error JSON, 退出码 3."""
    res = subprocess.run(
        [sys.executable, str(PARSE_PY), "/tmp/definitely-not-exist-xxx.txt"],
        capture_output=True, text=True,
    )
    assert res.returncode == 3
    err = json.loads(res.stdout)
    assert "不存在" in err["error"]


def test_no_args() -> None:
    """没 args → error, 退出码 2."""
    res = subprocess.run(
        [sys.executable, str(PARSE_PY)],
        capture_output=True, text=True,
    )
    assert res.returncode == 2


# ============================================================
# 关键防回归: 5/5 鸿波报"截断了不能用"
# ============================================================

def test_excel_huge_does_not_dump_full_data(tmp_path: Path, has_openpyxl: bool) -> None:
    """1000 行 Excel: preview 不应含全部数据 (那就是当年的 bug),
    应只含前 ~20 行 + 总行数提示."""
    from openpyxl import Workbook  # noqa: PLC0415
    f = tmp_path / "huge.xlsx"
    wb = Workbook()
    s = wb.active
    s.append(["id", "value"])
    for i in range(1000):
        # 用易识别的 marker 防 preview 误命中
        s.append([i, f"row_data_{i:04d}"])
    wb.save(f)

    r = run_parse(f)
    assert r["kind"] == "excel"
    # 前 20 行在 preview
    assert "row_data_0000" in r["preview_text"]
    # 第 500 行不在 preview (preview-only mode 核心)
    assert "row_data_0500" not in r["preview_text"]
    # 第 999 行不在
    assert "row_data_0999" not in r["preview_text"]
    # 提示用 execute_code
    assert "execute_code" in r["preview_text"]


# ============================================================
# BL-I4 audio + BL-I3.1 video — 5/8 ship
# ============================================================
#
# 真转写需要 ffmpeg + whisper-cli + ggml-small.bin 模型 (~466MB),
# CI 跑不到 — 这里只测 PARSERS 注册 + import + module 内可纯 Python 测的部分.
# 真音频转写测试用 manual_test_audio.sh (员工 mac 真录音 → run_parse 看输出).


def test_audio_extensions_registered() -> None:
    """所有支持的音频扩展都在 PARSERS dict, kind='audio'"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf
    expected = [".mp3", ".wav", ".m4a", ".flac", ".aac", ".ogg"]
    for ext in expected:
        assert ext in pf.PARSERS, f"audio 扩展 {ext} 没注册"
        kind, parser = pf.PARSERS[ext]
        assert kind == "audio"
        assert parser is pf.parse_audio_preview


def test_video_extensions_registered() -> None:
    """所有支持的视频扩展都在 PARSERS dict, kind='video'"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf
    expected = [".mp4", ".mov", ".m4v", ".mkv", ".webm"]
    for ext in expected:
        assert ext in pf.PARSERS, f"video 扩展 {ext} 没注册"
        kind, parser = pf.PARSERS[ext]
        assert kind == "video"
        assert parser is pf.parse_video_preview


def test_audio_video_in_full_text_extractors() -> None:
    """audio/video 在 _FULL_TEXT_EXTRACTORS 里, 大文件 (≥50KB 转写) 走 BM25 sidecar"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf
    assert "audio" in pf._FULL_TEXT_EXTRACTORS
    assert "video" in pf._FULL_TEXT_EXTRACTORS
    # 都指向同一个抽全文函数 (内部用 _transcribe_audio_to_text)
    assert pf._FULL_TEXT_EXTRACTORS["audio"] is pf.extract_full_text_audio
    assert pf._FULL_TEXT_EXTRACTORS["video"] is pf.extract_full_text_audio


def test_transcribe_missing_ffmpeg_raises_clear_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """ffmpeg 没装时报清楚错, 让员工知道装啥"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf
    monkeypatch.setattr(pf, "_find_executable", lambda name: None)
    with pytest.raises(RuntimeError, match=r"ffmpeg"):
        pf._transcribe_audio_to_text(Path("/tmp/fake.mp3"))


def test_transcribe_missing_whisper_raises_clear_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """ffmpeg 装了但 whisper-cli 没装"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf

    def fake_which(name: str) -> str | None:
        return "/usr/bin/ffmpeg" if name == "ffmpeg" else None

    monkeypatch.setattr(pf, "_find_executable", fake_which)
    with pytest.raises(RuntimeError, match=r"whisper-cli"):
        pf._transcribe_audio_to_text(Path("/tmp/fake.mp3"))


def test_transcribe_missing_model_raises_clear_error(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    """ffmpeg + whisper-cli 都装了, 但模型没下载"""
    sys.path.insert(0, str(THIS.parent))
    import parse_file as pf

    def fake_which(name: str) -> str | None:
        return f"/usr/bin/{name}"

    monkeypatch.setattr(pf, "_find_executable", fake_which)
    # 把 whisper 模型路径指到一个空 tmpdir
    monkeypatch.setattr(pf.Path, "home", lambda: tmp_path)
    with pytest.raises(RuntimeError, match=r"whisper.*模型"):
        pf._transcribe_audio_to_text(Path("/tmp/fake.mp3"))
