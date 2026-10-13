"""P54 只读代码免批: readonly_code.analyze 的判定 + 绕过尝试 + guard 包装。

原则是「只认得出来的才放行, 拿不准一律交给人」。这里两半:
  · 真实只读脚本 (10/13 微信里那几段) 必须被认出来 —— 否则免批没意义;
  · 各种绕过写法必须**全部**判有副作用 —— 漏一个就是开了口子。
"""
from __future__ import annotations

import textwrap

import pytest

import plugin_approval_ux as ux
import readonly_code as rc


def ro(code: str) -> bool:
    return rc.analyze(textwrap.dedent(code)).read_only


# ── 必须放行: 10/13 微信里真实出现的只读脚本 ───────────────────────────

REAL_1013 = [
    # 福富资质: 列目录 + 读 md + 正则抽字段
    """
    import os, json, re, collections
    base = os.path.expanduser('~/.catfish/wiki/entities')
    rows = []
    for fn in sorted(os.listdir(base)):
        if not fn.endswith('.md'):
            continue
        p = os.path.join(base, fn)
        t = open(p, encoding='utf-8').read()
        def field(k):
            m = re.search(r'^%s:\\s*(.*)$' % k, t, re.M)
            return m.group(1) if m else ''
        rows.append((fn, field('title')))
    print(json.dumps(rows, ensure_ascii=False)[:2000])
    """,
    """
    import os
    base=os.path.expanduser('~/.catfish/wiki')
    for rel in ['concepts/荣誉类剔除原则.md','concepts/知识库实持优先原则.md']:
        p=os.path.join(base,rel)
        print('####',rel, os.path.exists(p))
        if os.path.exists(p):
            print(open(p,encoding='utf-8').read()[:1500])
    """,
    """
    import os
    p=os.path.expanduser('~/.catfish/wiki/concepts/知识库全量口径.md')
    print(open(p,encoding='utf-8').read().split('---')[-1][:900])
    """,
    # 早先那段: ps 看进程
    """
    import subprocess, time, datetime
    out = subprocess.run(["ps", "-axo", "pid,etime,command"], capture_output=True, text=True).stdout
    for ln in out.splitlines():
        ll = ln.lower()
        if 'hermes' in ll:
            print(ln)
    """,
]


@pytest.mark.parametrize("code", REAL_1013)
def test_real_readonly_scripts_from_1013_are_exempt(code):
    v = rc.analyze(textwrap.dedent(code))
    assert v.read_only, v.effects


def test_other_common_readonly_patterns():
    assert ro("""
        from pathlib import Path
        import openpyxl, pandas as pd, glob
        for f in glob.glob('/Users/x/Downloads/*.xlsx'):
            wb = openpyxl.load_workbook(f, read_only=True)
            print(wb.sheetnames)
        df = pd.read_excel('a.xlsx'); print(df.head())
        print(Path('a.md').read_text())
        with open('b.csv', 'r', encoding='utf-8') as fh:
            print(fh.read())
        import sys; sys.stdout.write('ok')
    """)
    assert ro('import subprocess\nprint(subprocess.check_output(["git", "log", "--oneline", "-5"]))')
    assert ro('import subprocess\nsubprocess.run("ls -la ~/Downloads | grep xlsx", shell=True)')
    assert ro('from hermes_tools import read_file, search_files\nprint(read_file("a.md"))')


# ── 必须审批: 写 / 删 / 联网 / 跑命令 ──────────────────────────────────

SIDE_EFFECTS = {
    "写文件": "open('a.txt', 'w').write('x')",
    "追加": "open('a.txt', mode='a')",
    "Path 写": "from pathlib import Path\nPath('a').write_text('x')",
    "Path open w": "from pathlib import Path\nPath('a').open('w')",
    "删文件": "import os\nos.remove('a')",
    "删目录树": "import shutil\nshutil.rmtree('/tmp/x')",
    "改名": "import os\nos.rename('a', 'b')",
    "excel 保存": "import openpyxl\nwb = openpyxl.Workbook()\nwb.save('a.xlsx')",
    "pandas 导出": "import pandas as pd\npd.DataFrame().to_excel('a.xlsx')",
    "联网": "import requests\nrequests.get('https://x')",
    "urllib": "from urllib.request import urlopen\nurlopen('http://x')",
    "socket": "import socket",
    "rm 命令": "import subprocess\nsubprocess.run(['rm', '-rf', '/tmp/x'])",
    "os.system": "import os\nos.system('ls')",
    "非字面量命令": "import subprocess\ncmd = ['ls']\nsubprocess.run(cmd)",
    "find delete": "import subprocess\nsubprocess.run(['find', '.', '-delete'])",
    "sed -i": "import subprocess\nsubprocess.run(['sed', '-i', '', 's/a/b/', 'f'])",
    "git commit": "import subprocess\nsubprocess.run(['git', 'commit', '-m', 'x'])",
    "git -c 注入": "import subprocess\nsubprocess.run(['git', '-c', 'core.pager=sh', 'log'])",
    "sort -o": "import subprocess\nsubprocess.run(['sort', '-o', 'out', 'in'])",
    "uniq 写第二个文件": "import subprocess\nsubprocess.run(['uniq', 'in', 'out'])",
    "awk system": "import subprocess\nsubprocess.run(['awk', 'BEGIN{system(\"rm x\")}'])",
    "defaults write": "import subprocess\nsubprocess.run(['defaults', 'write', 'a', 'b', 'c'])",
    "python -c": "import subprocess\nsubprocess.run(['python3', '-c', 'print(1)'])",
    "hermes 写工具": "from hermes_tools import write_file\nwrite_file('a', 'b')",
    "hermes 终端": "import hermes_tools\nhermes_tools.terminal('rm -rf x')",
    "kill": "import os\nos.kill(1, 9)",
    "sqlite": "import sqlite3\nsqlite3.connect('a.db').execute('delete from t')",
    "不认识的模块": "import some_unknown_pkg",
    "解析不了": "def (",
}


@pytest.mark.parametrize("label", sorted(SIDE_EFFECTS))
def test_side_effects_need_approval(label):
    v = rc.analyze(SIDE_EFFECTS[label])
    assert not v.read_only, label
    assert v.effects, label


# ── 绕过尝试: 每一条都必须判「有副作用」 ────────────────────────────────

BYPASSES = {
    "shell 串命令": "import subprocess\nsubprocess.run('ls; rm -rf x', shell=True)",
    "shell 换行": "import subprocess\nsubprocess.run('ls\\nrm -rf x', shell=True)",
    "列表参数换行": "import subprocess\nsubprocess.run(['ls', 'a\\nrm'])",
    "shell 重定向": "import subprocess\nsubprocess.run('echo x > a', shell=True)",
    "命令替换": "import subprocess\nsubprocess.run('ls $(rm x)', shell=True)",
    "反引号": "import subprocess\nsubprocess.run('ls `rm x`', shell=True)",
    "getoutput": "import subprocess\nsubprocess.getoutput('rm -rf x')",
    "函数改名": "import subprocess\nr = subprocess.run\nr(['rm', 'x'])",
    "模块改名": "import subprocess\nsp = subprocess\nsp.run(['rm', 'x'])",
    "from 导入后改名": "from subprocess import run\nf = run\nf(['rm', 'x'])",
    "from 导入后传参": "from subprocess import run\nlist(map(run, [['rm', 'x']]))",
    "as 别名正常调用仍检查": "import subprocess as s\ns.run(['rm', 'x'])",
    "from 导入直接调用仍检查": "from subprocess import check_output as co\nco(['rm', 'x'])",
    "getattr 字面量": "import os\ngetattr(os, 'system')('rm x')",
    "getattr 拼接": "import os\ngetattr(os, 'sys' + 'tem')('rm x')",
    "getattr subprocess": "import subprocess\ngetattr(subprocess, 'run')(['rm', 'x'])",
    "sys.modules": "import sys\nsys.modules['os'].system('rm x')",
    "sys.modules subprocess": "import sys, subprocess\nsys.modules['subprocess'].run(['rm', 'x'])",
    "__import__": "__import__('os').system('rm x')",
    "eval": "eval('1')",
    "exec": "exec('import os')",
    "compile": "compile('x', 'f', 'exec')",
    "builtins 模块": "import builtins\nbuiltins.open('a', 'w')",
    "__builtins__": "__builtins__['open']('a', 'w')",
    "dunder 链": "().__class__.__base__.__subclasses__()",
    "vars": "import os\nvars(os)['system']('x')",
    "globals": "globals()['x'] = 1",
    "open 改名": "o = open\no('a', 'w')",
    "open 当参数": "list(map(open, ['a']))",
    "open 非字面量 mode": "m = 'w'\nopen('a', m)",
    "io.open 写": "import io\nio.open('a', 'w')",
    "codecs.open 写": "import codecs\ncodecs.open('a', 'w')",
    "io.open 改名": "import io\nf = io.open\nf('a', 'w')",
    "import *": "from os import *\nsystem('x')",
    "os 写函数直接 from": "from os import remove\nremove('a')",
    "os 写函数 as": "from os import system as s\ns('x')",
    "os 别名": "import os as o\no.remove('a')",
    "importlib": "import importlib\nimportlib.import_module('os').system('x')",
    "相对导入": "from . import x",
    "kwargs 展开": "import subprocess\nkw={'shell': True}\nsubprocess.run('ls; rm x', **kw)",
    "stdout 写文件": "import subprocess\nsubprocess.run(['ls'], stdout=open('a', 'w'))",
    "stdout 变量": "import subprocess\nf = None\nsubprocess.run(['ls'], stdout=f)",
    "executable": "import subprocess\nsubprocess.run(['ls'], executable='/bin/rm')",
    "rg --pre": "import subprocess\nsubprocess.run(['rg', '--pre', 'sh', 'x'])",
    "git --output": "import subprocess\nsubprocess.run(['git', 'log', '--output=/tmp/x'])",
    "git ext-diff": "import subprocess\nsubprocess.run(['git', 'diff', '--ext-diff'])",
    "git config 写": "import subprocess\nsubprocess.run(['git', 'config', 'user.name', 'x'])",
    "awk -f": "import subprocess\nsubprocess.run(['awk', '-f', 'prog.awk'])",
    "sed w 命令": "import subprocess\nsubprocess.run(['sed', 'w out', 'in'])",
    "sed -f": "import subprocess\nsubprocess.run(['sed', '-f', 'script.sed', 'in'])",
    "tree -o": "import subprocess\nsubprocess.run(['tree', '-o', 'out'])",
    "date 改时间": "import subprocess\nsubprocess.run(['date', '0101000026'])",
    "hostname 改名": "import subprocess\nsubprocess.run(['hostname', 'evil'])",
    "绝对路径伪装": "import subprocess\nsubprocess.run(['/tmp/ls'])" ,
    "env 跑别的": "import subprocess\nsubprocess.run(['env', 'rm', 'x'])",
    "pickle": "import pickle\npickle.loads(b'')",
    "ctypes": "import ctypes",
    "threading": "import threading",
    "zip 写": "import zipfile\nzipfile.ZipFile('a.zip', 'w')",
    "extractall": "import zipfile\nzipfile.ZipFile('a.zip').extractall('/')",
    "json.dump 到文件": "import json\njson.dump({}, fh)",
    "np.save": "import numpy as np\nnp.save('a', 1)",
}


@pytest.mark.parametrize("label", sorted(BYPASSES))
def test_bypass_attempts_are_caught(label):
    v = rc.analyze(BYPASSES[label])
    assert not v.read_only, f"{label}: 被当成只读了 —— 这是个口子"


def test_absolute_path_to_real_program_is_still_checked():
    # /bin/ls 本身只读, 放行没问题; /bin/rm 必须拦
    assert ro("import subprocess\nsubprocess.run(['/bin/ls', '-la'])")
    assert not ro("import subprocess\nsubprocess.run(['/bin/rm', 'x'])")


# ── 大白话描述 ────────────────────────────────────────────────────────

def test_describe_prefers_comment_then_files():
    v = rc.analyze("# 读福富的资质条目\nimport os\nprint(open('a.md').read())")
    assert rc.describe(v) == "读福富的资质条目"
    v = rc.analyze(textwrap.dedent(REAL_1013[1]))
    assert rc.describe(v).startswith("读取 荣誉类剔除原则.md、知识库实持优先原则.md")
    v = rc.analyze("import os\nprint(os.listdir('.'))")
    assert rc.describe(v) == "查看文件夹里有哪些文件"
    v = rc.analyze(textwrap.dedent(REAL_1013[3]))
    assert "查看正在运行的程序" in rc.describe(v)


def test_prompt_shows_no_code_and_explains_why():
    cmd = "execute_code <<'PY'\nimport openpyxl\nwb = openpyxl.Workbook()\nwb.save('资质台账.xlsx')\nPY"
    msg = ux.format_approval_prompt(cmd, "x")
    assert "要做的事: 处理 资质台账.xlsx" in msg
    assert "原因: 它会写入 / 删除 / 改动文件" in msg
    assert "openpyxl" not in msg and "```" not in msg and "wb.save" not in msg
    assert "回复 1 执行 · 2 不要执行" in msg


# ── guard 包装 ────────────────────────────────────────────────────────

def _guard():
    calls = []

    def orig(code, env_type, has_host_access=False):
        calls.append(code)
        return {"approved": False, "message": "need approval"}

    return ux.wrap_execute_code_guard(orig, rc), calls


def test_guard_auto_approves_readonly_only():
    g, calls = _guard()
    assert g("import os\nprint(os.listdir('.'))", "local")["approved"] is True
    assert calls == []
    r = g("import os\nos.remove('a')", "local", True)
    assert r == {"approved": False, "message": "need approval"}
    assert calls == ["import os\nos.remove('a')"]


def test_guard_falls_back_to_approval_when_analysis_breaks():
    class Broken:
        @staticmethod
        def analyze(code):
            raise RuntimeError("boom")

    calls = []
    g = ux.wrap_execute_code_guard(lambda c, e, h=False: calls.append(c) or {"approved": False}, Broken)
    assert g("print(1)", "local") == {"approved": False}
    assert calls == ["print(1)"]
