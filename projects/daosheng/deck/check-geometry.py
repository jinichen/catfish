#!/usr/bin/env python3
"""PPT 版面自检 —— build 后必跑，越界不算完成。

# 为啥要这个

7/28 鸿波在 v30 发现 P4 的方块跑出底板，说"我再去手动改很麻烦"。
根因是布局宽度写死常数（卡片起点 2.60，可用宽却按 10.50 算），
靠人眼在 24 页里挑这种 0.4 英寸的溢出不现实，一次改完下次还会犯。

所以把它变成机器检查：读生成的 pptx，算每个形状的四边，超版心即报错。

# 版心

    x  0.60 – 12.70    各页底板统一 x=0.6 w=12.1
    y  0.50 –  7.45    页脚基线 7.15，其下留 0.3 余量

# 排除项

满幅装饰条（x≈0 且几乎与画布同宽）是有意为之，不算越界。
封面顶部那道金色横条就是这种。

# 用法

    node build-vNN.js && python3 check-geometry.py daosheng-vNN-smart-mfg.pptx

退出码非 0 表示有越界，可直接串在 build 后面。
"""
from __future__ import annotations

import sys
from collections import defaultdict

try:
    from pptx import Presentation
except ImportError:
    sys.exit("需要 python-pptx: pip install python-pptx --break-system-packages")

EMU = 914400.0

# 版心边界（英寸）
LEFT, RIGHT = 0.60, 12.70
TOP, BOTTOM = 0.50, 7.45
TOL = 0.02          # 容差，避免浮点零头误报


def is_full_bleed(x: float, w: float, slide_w: float) -> bool:
    """满幅装饰条：贴左边缘且几乎横贯画布 —— 有意为之，跳过。"""
    return x <= 0.05 and w >= slide_w - 0.1


def label(shape) -> str:
    if getattr(shape, "has_table", False):
        try:
            head = " | ".join(c.text.strip() for c in shape.table.rows[0].cells)[:30]
            return f"表格[{len(shape.table.rows)}行]「{head}」"
        except Exception:
            return "表格"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return shape.text_frame.text.strip().replace("\n", " / ")[:34]
    return f"<{shape.shape_type}>"


def is_filled(shape) -> bool:
    """有实心填充的形状才会遮挡别的内容；透明文本框不算。"""
    try:
        return shape.fill.type == 1        # MSO_FILL.SOLID
    except Exception:
        return False


def box(shape) -> tuple[float, float, float, float]:
    x, y = shape.left / EMU, (shape.top or 0) / EMU
    return x, y, x + shape.width / EMU, y + (shape.height or 0) / EMU


def contains(a, b, tol: float = 0.02) -> bool:
    """a 完全包住 b —— 卡片套色条、底板套卡片都属于这种，合法。"""
    return (a[0] <= b[0] + tol and a[1] <= b[1] + tol
            and a[2] >= b[2] - tol and a[3] >= b[3] - tol)


def overlap_area(a, b) -> float:
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def table_box(sh) -> tuple[float, float, float, float] | None:
    """表格的真实占位。

    ★ 8/6 补。此前这个脚本**只比对 addShape 的实心块，完全不看 addTable**，
    于是「表格最后一行被底部提示框压住」这类问题一次都抓不到 ——
    上面那句注释写着「典型命中：表格最后一行被底部结论框压住」，
    其实压根检不出来，v2 的 P9 / P15 都是这么漏过去的，最后靠渲染 JPEG 用眼睛看出来。

    高度**以各行 trHeight 之和为准，不用 graphicFrame 的 height**：
    pptxgenjs 给 graphicFrame 写的是个占位值（实测恒为 1.00″，与真实行数无关），
    2 行的表和 5 行的表都写 1.00，拿它比对会一边漏报一边误报。

    注意这仍是**下界**：rowH 装不下文字时 PowerPoint / LibreOffice 会自动撑高，
    而撑高不写回 XML。所以表底和下方提示框之间要留余量，卡得刚好等于没检查。
    """
    if not getattr(sh, "has_table", False):
        return None
    x, y = sh.left / EMU, (sh.top or 0) / EMU
    w = sh.width / EMU
    try:
        summed = sum((r.height or 0) for r in sh.table.rows) / EMU
    except Exception:
        summed = 0.0
    h = summed if summed > 0 else (sh.height or 0) / EMU
    return x, y, x + w, y + h


def find_occlusions(slide, slide_w: float) -> list[tuple[float, str, str]]:
    """找互相遮挡的实心块 —— 现在把表格也算进来。

    只看**两个都有填充**的形状：透明文本框盖不住东西。
    互相包含的跳过（那是正常的层叠：底板 → 卡片 → 色条）。
    """
    blocks = []
    for sh in slide.shapes:
        if sh.left is None or sh.width is None:
            continue
        tb = table_box(sh)
        if tb is not None:                      # 表格：一律参与比对
            if not is_full_bleed(tb[0], tb[2] - tb[0], slide_w):
                blocks.append((tb, sh))
            continue
        if not is_filled(sh):
            continue
        b = box(sh)
        if is_full_bleed(b[0], b[2] - b[0], slide_w):
            continue
        blocks.append((b, sh))

    hits = []
    for i in range(len(blocks)):
        for j in range(i + 1, len(blocks)):
            a, sa = blocks[i]
            b, sb = blocks[j]
            # ★ 8/6 第二个坑：contains 豁免只对「底板 → 卡片 → 色条」这类层叠成立。
            # 表格不可能合法地「包住」别的东西 —— 而一个提示框若整个落在表格的
            # 包围盒里，正是我们要抓的那种「表格最后一行被压住」。
            # P19 就是这么漏掉的：表底 5.65，提示框 5.05–5.55，完全落在表内 → 被豁免。
            involves_table = (getattr(sa, "has_table", False)
                              or getattr(sb, "has_table", False))
            if not involves_table and (contains(a, b) or contains(b, a)):
                continue
            area = overlap_area(a, b)
            if area > 0.015:                       # 约 0.12″ × 0.12″ 以上才算
                hits.append((round(area, 3), label(sa), label(sb)))
    return sorted(hits, reverse=True)


def main() -> int:
    path = sys.argv[1] if len(sys.argv) > 1 else "daosheng-v31-smart-mfg.pptx"
    prs = Presentation(path)
    slide_w = prs.slide_width / EMU
    slide_h = prs.slide_height / EMU

    issues: dict[int, list[tuple[str, float, str]]] = defaultdict(list)

    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is None or sh.width is None:
                continue
            x, y = sh.left / EMU, (sh.top or 0) / EMU
            w, h = sh.width / EMU, (sh.height or 0) / EMU

            if is_full_bleed(x, w, slide_w):
                continue

            if x + w > RIGHT + TOL:
                issues[idx].append(("右", round(x + w - RIGHT, 2), label(sh)))
            if x < LEFT - TOL:
                issues[idx].append(("左", round(LEFT - x, 2), label(sh)))
            if y + h > BOTTOM + TOL:
                issues[idx].append(("下", round(y + h - BOTTOM, 2), label(sh)))
            if y < TOP - TOL and y > 0.01:      # y=0 多为满幅背景/色条
                issues[idx].append(("上", round(TOP - y, 2), label(sh)))

    # ── 遮挡检查 ──
    occl: dict[int, list[tuple[float, str, str]]] = {}
    for idx, slide in enumerate(prs.slides, 1):
        hits = find_occlusions(slide, slide_w)
        if hits:
            occl[idx] = hits

    print(f"画布 {slide_w:.2f} × {slide_h:.2f} 英寸")
    print(f"版心 x {LEFT}–{RIGHT} · y {TOP}–{BOTTOM} · 容差 {TOL}\n")

    rc = 0

    print("【一】越界")
    if not issues:
        print(f"  ✓ {len(prs.slides)} 页全部在版心内")
    else:
        rc = 1
        total = sum(len(v) for v in issues.values())
        print(f"  ✗ {total} 处，分布在 {len(issues)} 页：")
        for pg in sorted(issues):
            worst = max(d for _, d, _ in issues[pg])
            print(f"    p{pg:<3d} {len(issues[pg])} 处 · 最大超出 {worst}\"")
            for side, delta, text in sorted(issues[pg], key=lambda t: -t[1])[:8]:
                print(f"          {side}越 {delta:>5.2f}\"   {text}")

    # ── 页脚保护区：正文不得压到页脚上 ──
    # 遮挡检查只看实心块，透明文本框盖住页脚它抓不到 ——
    # v32 的 P13 底部注释就是这样压在页脚上没被发现的。
    FOOTER_TOP = 7.10        # 页脚基线 7.15，上方留 0.05 余量
    intrude: dict[int, list[tuple[float, str]]] = {}
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is None or sh.width is None:
                continue
            y = (sh.top or 0) / EMU
            bot = y + (sh.height or 0) / EMU
            if y >= 7.14:                    # 页脚自身
                continue
            if bot > FOOTER_TOP + TOL:
                intrude.setdefault(idx, []).append(
                    (round(bot - FOOTER_TOP, 2), label(sh)))

    print("\n【二】遮挡（两个实心块互相压住，且非包含关系）")
    if not occl:
        print(f"  ✓ 无遮挡")
    else:
        rc = 1
        total = sum(len(v) for v in occl.values())
        print(f"  ✗ {total} 处，分布在 {len(occl)} 页：")
        for pg in sorted(occl):
            print(f"    p{pg:<3d} {len(occl[pg])} 处 · 最大重叠 {occl[pg][0][0]} 平方英寸")
            for area, a, b in occl[pg][:5]:
                print(f"          {area:>5.3f}   「{a or '—'}」 ✕ 「{b or '—'}」")

    print(f"\n【三】压页脚（正文底边越过 {FOOTER_TOP}″）")
    if not intrude:
        print("  ✓ 无侵入")
    else:
        rc = 1
        total = sum(len(v) for v in intrude.values())
        print(f"  ✗ {total} 处，分布在 {len(intrude)} 页：")
        for pg in sorted(intrude):
            for delta, text in sorted(intrude[pg], reverse=True)[:4]:
                print(f"    p{pg:<3d} 超 {delta:>5.2f}″   {text}")

    return rc


if __name__ == "__main__":
    sys.exit(main())
